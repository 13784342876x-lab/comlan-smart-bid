import json
from pathlib import Path

from docx import Document
from openpyxl import Workbook

from app.technical_outline_builder import build_technical_outline, render_technical_outline_markdown


def add_paragraphs(document: Document, values: list[str]) -> None:
    for value in values:
        document.add_paragraph(value)


def write_selected_package(directory: Path) -> None:
    directory.mkdir(parents=True, exist_ok=True)

    procurement = Document()
    add_paragraphs(
        procurement,
        [
            "第六章 投标文件格式",
            "三、技术文件",
            "1. 技术文件组成",
            "*投标人必须提交技术响应文件、技术参数响应表和技术附件。",
            "2. 技术偏差表",
            "投标人应逐项填写技术偏差表，不得遗漏。",
            "二、商务文件",
            "3. 法定代表人授权书",
            "投标人应提交法定代表人授权委托书。",
            "第七章 价格文件",
            "1. 价格评分",
            "最高限价为100万元，报价税率为6%。",
        ],
    )
    procurement.save(directory / "采购文件.docx")

    specification = Document()
    long_clause = (
        "（1）投标人应提供满足本规范书中要求的全部技术研究内容、研究技术指标、培训和服务；"
        "项目实施过程中必须建立数据可信、权限控制和安全审计机制，并提交完整的测试记录、技术文档和验收资料；"
        "未经采购人书面确认不得删减研究任务。"
    )
    add_paragraphs(
        specification,
        [
            "第五章 技术规范书",
            "3.2 服务内容",
            long_clause,
            "（2）*系统必须兼容现有数据接口，不得改变原有业务数据。",
            "3.3 验收要求",
            "项目成果应通过功能测试、性能测试和安全测试，验收时应提交测试报告。",
            "附件1 技术参数表",
        ],
    )
    table = specification.add_table(rows=3, cols=4)
    for row, values in zip(
        table.rows,
        [
            ["序号", "技术参数", "要求", "强制标记"],
            ["1", "权限控制", "应支持分级授权", "*"],
            ["2", "安全审计", "审计记录不得删除", "★"],
        ],
    ):
        for cell, value in zip(row.cells, values):
            cell.text = value
    specification.save(directory / "技术规范书.docx")

    score = Document()
    add_paragraphs(score, ["第三章 评审办法", "技术评分标准"])
    tech_table = score.add_table(rows=4, cols=4)
    for row, values in zip(
        tech_table.rows,
        [
            ["评分项", "分值", "评分标准", "证明材料要求"],
            ["技术方案", "10", "优：9-10分；良：7-8分；一般：6分。", "提供技术响应方案"],
            ["项目团队", "5", "项目负责人和团队人员配置完整得5分。", "提供人员简历和证书"],
            ["业绩及资源实力", "4", "每提供1项技术服务业绩得2分，最高4分。", "提供合同和验收证明"],
        ],
    ):
        for cell, value in zip(row.cells, values):
            cell.text = value
    score.add_paragraph("商务评分标准")
    business_table = score.add_table(rows=2, cols=3)
    for row, values in zip(
        business_table.rows,
        [["商务评分项", "分值", "标准"], ["信用报告", "5", "信用等级高得5分"]],
    ):
        for cell, value in zip(row.cells, values):
            cell.text = value
    score.add_paragraph("价格评分标准")
    price_table = score.add_table(rows=2, cols=3)
    for row, values in zip(
        price_table.rows,
        [["价格评分项", "分值", "标准"], ["评标价", "30", "按基准价计算报价得分"]],
    ):
        for cell, value in zip(row.cells, values):
            cell.text = value
    score.save(directory / "评标办法.docx")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "技术偏差表"
    sheet.append(["序号", "条款号", "招标要求", "投标响应", "偏差说明", "报价"])
    sheet.append([1, "3.2（1）", "必须提供全部技术研究内容", "", "", "不得在技术表填写价格"])
    workbook.save(directory / "技术偏差表.xlsx")

    (directory / "项目需求书.txt").write_text(
        "2.1 采购范围\n投标人应提供三年技术支持服务。\n2.2 服务期限\n服务期限为自合同签订之日起3年。",
        encoding="utf-8",
    )
    (directory / "技术附件A.txt").write_text("服务期限为自合同签订之日起3年。", encoding="utf-8")
    (directory / "技术附件B.txt").write_text("服务期限为自合同签订之日起 3 年。", encoding="utf-8")

    manual = Document()
    add_paragraphs(
        manual,
        ["投标工具操作手册", "投标人登录电子招标投标交易平台后使用CA证书解密并递交投标文件。"],
    )
    manual.save(directory / "投标工具操作手册.docx")


def load_outputs(result: dict) -> tuple[dict, str]:
    data = json.loads(Path(result["files"]["json"]).read_text(encoding="utf-8"))
    markdown = Path(result["files"]["markdown"]).read_text(encoding="utf-8")
    return data, markdown


def all_points(data: dict) -> list[dict]:
    return [point for chapter in data["chapters"] for point in chapter["response_points"]]


def test_selected_directory_builds_clean_technical_outline(tmp_path: Path) -> None:
    selected = tmp_path / "0425139003001-6582"
    write_selected_package(selected)
    result = build_technical_outline(selected, output_dir=tmp_path / "outputs", enable_page_mapping=False)
    data, markdown = load_outputs(result)
    points = all_points(data)
    text = "\n".join(point["original_text"] for point in points)

    assert data["schema_version"] == "technical_response_outline_v2"
    assert data["scope"] == "technical_only"
    assert data["source_mode"] == "manually_selected_files"
    assert "*投标人必须提交技术响应文件、技术参数响应表和技术附件。" in text
    assert "投标人应逐项填写技术偏差表，不得遗漏。" in text
    assert "未经采购人书面确认不得删减研究任务。" in text
    assert "项目成果应通过功能测试、性能测试和安全测试" in text
    assert "优：9-10分；良：7-8分；一般：6分。" in text
    assert "提供人员简历和证书" in text
    assert "提供合同和验收证明" in text
    assert "投标人应提供三年技术支持服务。" in text
    assert not any(point["original_text"] == "1.3对投标人的技术要求" for point in points)
    assert any("3.2 服务内容" in point.get("section_path", []) for point in points)
    assert any("3.3 验收要求" in point.get("section_path", []) for point in points)

    prohibited = [
        "最高限价",
        "价格评分",
        "报价税率",
        "信用报告",
        "法定代表人授权",
        "投标保证金",
        "平台解密",
        "CA证书",
        "当前知识库状态",
        "候选人员",
        "推荐部门",
        "缺失材料",
        "材料准备清单",
        "technical_solution_inputs",
        "project_leader_resume",
        "company_contract",
        "评分项行动分析",
        "建议写作结构",
        "评分关注",
    ]
    rendered = json.dumps(data, ensure_ascii=False) + markdown
    assert not any(term in rendered for term in prohibited)
    assert render_technical_outline_markdown(data) == markdown
    assert not list((tmp_path / "outputs" / selected.name / "outline").glob("*debug*"))


def test_original_text_numbering_markers_order_and_long_clause_are_preserved(tmp_path: Path) -> None:
    selected = tmp_path / "selected"
    write_selected_package(selected)
    data, _ = load_outputs(build_technical_outline(selected, output_dir=tmp_path / "out", enable_page_mapping=False))
    points = all_points(data)
    texts = [point["original_text"] for point in points]
    joined = "\n".join(texts)

    long_clause = next(text for text in texts if "全部技术研究内容" in text)
    assert long_clause.startswith("（1）投标人应提供")
    assert "研究技术指标、培训和服务" in long_clause
    assert "必须建立数据可信、权限控制和安全审计机制" in long_clause
    assert "不得删减研究任务" in long_clause
    assert "*系统必须兼容现有数据接口，不得改变原有业务数据。" in joined
    assert any(point["mandatory_marker"] == "*" for point in points if "系统必须兼容" in point["original_text"])
    section_paths = [point.get("section_path", []) for point in points]
    assert next(i for i, path in enumerate(section_paths) if "3.2 服务内容" in path) < next(
        i for i, path in enumerate(section_paths) if "3.3 验收要求" in path
    )


def test_tables_scores_and_excel_locations_are_preserved(tmp_path: Path) -> None:
    selected = tmp_path / "selected"
    write_selected_package(selected)
    data, markdown = load_outputs(build_technical_outline(selected, output_dir=tmp_path / "out", enable_page_mapping=False))
    tables = [point for point in all_points(data) if point["content_type"] == "table"]
    table_text = "\n".join(point["original_text"] for point in tables)

    assert "序号\t技术参数\t要求\t强制标记" in table_text
    assert "1\t权限控制\t应支持分级授权\t*" in table_text
    assert "2\t安全审计\t审计记录不得删除\t★" in table_text
    assert "评分项\t分值\t评分标准\t证明材料要求" in table_text
    assert "技术方案\t10\t优：9-10分；良：7-8分；一般：6分。" in table_text
    assert "项目团队\t5" in table_text and "业绩及资源实力\t4" in table_text
    assert "商务评分项" not in table_text and "价格评分项" not in table_text
    deviation = next(point for point in tables if any(ref["sheet_name"] == "技术偏差表" for ref in point["source_refs"]))
    assert deviation["table"]["rows"][0] == ["序号", "条款号", "招标要求", "投标响应", "偏差说明"]
    assert deviation["source_refs"][0]["row_or_cell"] == "A1:E2"
    assert "| 评分项 | 分值 | 评分标准 | 证明材料要求 |" in markdown


def test_every_point_has_specific_source_and_exact_duplicates_only_are_merged(tmp_path: Path) -> None:
    selected = tmp_path / "selected"
    write_selected_package(selected)
    data, _ = load_outputs(build_technical_outline(selected, output_dir=tmp_path / "out", enable_page_mapping=False))
    points = all_points(data)
    for point in points:
        assert point["source_refs"]
        for source in point["source_refs"]:
            assert source["file_name"]
            assert source["clause_path"] or source["row_or_cell"] or source["page"]

    service_period = [point for point in points if "服务期限为自合同签订之日起" in point["original_text"]]
    assert len(service_period) == 2
    attachment_period = next(point for point in service_period if point["original_text"].startswith("服务期限"))
    assert len(attachment_period["source_refs"]) == 2
    assert any(point["original_text"].startswith("2.2 服务期限") for point in service_period)
    assert any("投标人应提供三年技术支持服务" in point["original_text"] for point in points)
    assert data["statistics"]["exact_duplicate_count"] >= 1


def test_multiple_unclear_technical_score_tables_are_not_mixed(tmp_path: Path) -> None:
    selected = tmp_path / "selected"
    selected.mkdir()
    document = Document()
    for title in ["服务类评分表A", "服务类评分表B"]:
        document.add_paragraph(title)
        table = document.add_table(rows=2, cols=3)
        for row, values in zip(table.rows, [["评分项", "分值", "评分标准"], ["技术方案", "10", "优：9-10分"]]):
            for cell, value in zip(row.cells, values):
                cell.text = value
    document.save(selected / "评标办法.docx")

    result = build_technical_outline(selected, output_dir=tmp_path / "out", enable_page_mapping=False)
    data, _ = load_outputs(result)
    assert not any(point["content_type"] == "table" for point in all_points(data))
    assert any("多套可能适用的技术评分表" in warning for warning in data["build_warnings"])
