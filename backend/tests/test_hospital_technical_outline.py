import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.document_parser import parse_docx
from app.document_structure_normalizer import normalize_document_structure
from app.technical_outline_builder import build_technical_outline


def _add(document: Document, text: str, *, style: str | None = None, bold: bool = False, center: bool = False):
    paragraph = document.add_paragraph(style=style)
    run = paragraph.add_run(text)
    run.bold = bold
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return paragraph


def write_hospital_fixture(input_dir: Path) -> Path:
    input_dir.mkdir(parents=True)
    document = Document()
    _add(document, "第二部分 招标项目要求", style="Subtitle")
    _add(document, "一、商务要求")
    _add(document, "（一）报价要求")
    _add(document, "投标报价含税，付款比例按合同执行，投标保证金为零。")
    _add(document, "（二）服务要求")
    _add(document, "1. 提供所投产品至少5年的免费保修，终身维修。")
    _add(document, "2. 保修期内免费更换零配件，7×24小时技术响应，24小时内维修工程师到达维修现场。")
    _add(document, "3. 提供所投产品制造商服务机构情况。")
    _add(document, "4. 提供现场技术培训。")
    _add(document, "（三）交货要求")
    _add(document, "1. 交货期：")
    _add(document, "货到时间：自签订合同之日起30日内。")
    _add(document, "安装完成：货到之日起10日内。")
    _add(document, "2. 提供制造商完整的随机技术资料。")
    _add(document, "（六）验收方法及标准")
    _add(document, "应按照国家标准、行业标准逐项验收并出具验收书。")
    _add(document, "二、技术要求")
    _add(document, "（一）投标人须承诺所投产品和服务符合相关强制性规定。")
    _add(document, "（二）投标文件中应提供详细技术资料。")
    _add(document, "（三）投标文件中应提供产品性能质量证明材料。")
    _add(document, "（四）投标文件中应提供产品制造商能力证明材料。")
    _add(document, "（五）投标文件中应提供绿色供应链管理说明。")
    _add(document, "（六）具体需求详见本部分项目需求书。")
    _add(document, "三、评审因素及评标标准")
    score = document.add_table(rows=8, cols=4)
    score.cell(0, 0).merge(score.cell(0, 2)).text = "第一部分价格（30分）"
    score.cell(0, 3).text = "分值"
    for cell, value in zip(score.rows[1].cells, ["1", "价格", "投标报价得分按基准价计算", "30分"]):
        cell.text = value
    score.cell(2, 0).merge(score.cell(2, 2)).text = "第二部分客观分（40分）"
    score.cell(2, 3).text = "分值"
    for cell, value in zip(score.rows[3].cells, ["1", "环境标志产品", "提供环境标志证明得分", "10分"]):
        cell.text = value
    score.cell(4, 0).merge(score.cell(4, 2)).text = "第三部分主观分（30分）"
    score.cell(4, 3).text = "分值"
    for cell, value in zip(score.rows[5].cells, ["1", "实施部署评价", "方案完整得30分", "30分"]):
        cell.text = value
    score.cell(6, 0).merge(score.cell(6, 2)).text = "合计"
    score.cell(6, 3).text = "100分"
    score.cell(7, 0).merge(score.cell(7, 3)).text = "注：本项所称“瑕疵”是指内容缺项。"
    _add(document, "四、投标文件内容要求")
    _add(document, "（一）投标人须按照投标文件编制要求编制技术投标文件。")
    _add(document, "五、采购人职责")
    _add(document, "采购人负责组织内部审批和安排经办人员。")

    _add(document, "第三部分 投标须知", style="Subtitle")
    _add(document, "15. 投标文件格式", bold=True)
    _add(document, "18.2 上述文件可以是文字资料、图纸或数据，并须提供：")
    _add(document, "（3）逐条评议技术规格并完整填写《技术要求点对点应答表》。")
    _add(document, "18.3 投标文件中设备性能指标应达到或优于招标文件所列技术指标，投标应答必须列出具体数值或内容。")

    _add(document, "项目需求书", style="Subtitle")
    _add(document, "一、项目背景")
    _add(document, "为提升医院信息化水平，采购国产化超融合系统1套。")
    _add(document, "二、采购清单")
    parameters = document.add_table(rows=4, cols=4)
    for cell, value in zip(parameters.rows[0].cells, ["序号", "标的名称", "技术参数", "数量"]):
        cell.text = value
    parameters.cell(1, 0).merge(parameters.cell(2, 0)).text = "1"
    parameters.cell(1, 1).merge(parameters.cell(2, 1)).text = "超融合系统"
    parameters.cell(1, 2).text = "★支持数据高可用"
    parameters.cell(1, 3).merge(parameters.cell(2, 3)).text = "1套"
    parameters.cell(2, 2).text = "支持安全审计"
    for cell, value in zip(parameters.rows[3].cells, ["2", "超融合节点", "★支持横向扩展", "8台"]):
        cell.text = value
    _add(document, "注：")
    _add(document, "加注“★”号条款为实质性条款，不得出现负偏离。")
    _add(document, "加注“▲”号的产品为核心产品。")
    _add(document, "加注“■”号的产品属于节能产品政府采购强制采购的产品。")

    _add(document, "第五部分 投标文件格式", style="Subtitle")
    _add(document, "附件6-2")
    _add(document, "技术要求点对点应答表", bold=True, center=True)
    response = document.add_table(rows=3, cols=5)
    response.cell(0, 0).merge(response.cell(0, 4)).text = "招标文件第二部分技术要求"
    for cell, value in zip(response.rows[1].cells, ["序号", "招标要求", "投标应答", "偏离说明", "技术支撑材料页码"]):
        cell.text = value
    for cell, value in zip(response.rows[2].cells, ["1", "项目需求书要求", "填写具体参数", "填写偏离", "页码"]):
        cell.text = value
    _add(document, "1. 招标要求指招标文件中规定的具体要求，投标应答应填写具体内容。")
    _add(document, "2. 不得仅填写“符合”“满足”，应填写具体参数或具体内容。")

    path = input_dir / "医院招标文件.docx"
    document.save(path)

    specification = Document()
    _add(specification, "第五章 技术规范书", style="Subtitle")
    _add(specification, "附件1 技术参数表", bold=True, center=True)
    plain = specification.add_table(rows=2, cols=3)
    for cell, value in zip(plain.rows[0].cells, ["序号", "参数", "要求"]):
        cell.text = value
    for cell, value in zip(plain.rows[1].cells, ["1", "兼容性", "应兼容现有系统"]):
        cell.text = value
    specification.save(input_dir / "技术规范书.docx")
    return path


def load_result(tmp_path: Path) -> tuple[dict, str, Path]:
    input_dir = tmp_path / "input"
    document_path = write_hospital_fixture(input_dir)
    result = build_technical_outline(input_dir, output_dir=tmp_path / "outputs", enable_page_mapping=False)
    data = json.loads(Path(result["files"]["json"]).read_text(encoding="utf-8"))
    markdown = Path(result["files"]["markdown"]).read_text(encoding="utf-8")
    return data, markdown, document_path


def test_normalizer_recovers_normal_style_peer_headings(tmp_path: Path) -> None:
    _, _, path = load_result(tmp_path)
    parsed = parse_docx(path, {"id": "fixture", "file_ext": "docx", "extract_assets": False})
    blocks = normalize_document_structure(parsed["blocks"])
    by_text = {block.get("text"): block for block in blocks if block.get("block_type") == "paragraph"}

    assert by_text["二、技术要求"]["is_heading"] is True
    assert by_text["二、技术要求"]["heading_level"] == by_text["一、商务要求"]["heading_level"] == 2
    assert by_text["二、技术要求"]["heading_source"] == "rule_inferred"
    assert by_text["三、评审因素及评标标准"]["heading_level"] == 2
    assert by_text["项目需求书"]["heading_level"] == 1
    assert by_text["二、技术要求"]["parent_path"] == ["第二部分 招标项目要求"]


def test_outline_hierarchy_scope_and_subject_filtering(tmp_path: Path) -> None:
    data, _, _ = load_result(tmp_path)
    titles = [chapter["title"] for chapter in data["chapters"]]
    text = "\n".join(point["original_text"] for chapter in data["chapters"] for point in chapter["response_points"])

    assert titles[:5] == ["技术履约相关要求", "总体技术要求", "项目需求书", "技术相关评分标准", "技术投标文件格式与应答要求"]
    assert len(data["chapters"]) <= 6
    assert "至少5年的免费保修" in text and "现场技术培训" in text
    assert "安装完成" in text and "随机技术资料" in text and "验收书" in text
    assert "投标报价含税" not in text and "付款比例" not in text and "投标保证金" not in text
    assert "采购人负责组织内部审批" not in text
    assert "★支持数据高可用" in text and "★支持横向扩展" in text


def test_merged_cells_are_anchored_once_and_rendered_as_html(tmp_path: Path) -> None:
    data, markdown, _ = load_result(tmp_path)
    score = next(
        point
        for chapter in data["chapters"]
        if chapter["title"] == "技术相关评分标准"
        for point in chapter["response_points"]
        if point["content_type"] == "table"
    )
    cells = [cell for row in score["table"]["rows"] for cell in row["cells"]]

    assert score["table"]["has_merged_cells"] is True
    assert any(cell["text"] == "合计" and cell["colspan"] == 3 for cell in cells)
    assert sum(cell["text"] == "合计" for cell in cells) == 1
    assert sum("本项所称“瑕疵”" in cell["text"] for cell in cells) == 1
    assert '<td colspan="3">合计</td>' in markdown
    assert '<td colspan="4">注：本项所称“瑕疵”是指内容缺项。</td>' in markdown
    assert "合计 | 合计 | 合计" not in markdown


def test_tables_remain_whole_and_use_the_right_renderer(tmp_path: Path) -> None:
    data, markdown, _ = load_result(tmp_path)
    tables = [point for chapter in data["chapters"] for point in chapter["response_points"] if point["content_type"] == "table"]

    assert len(tables) == 4
    assert any(point["table"]["has_merged_cells"] for point in tables)
    assert any(not point["table"]["has_merged_cells"] for point in tables)
    assert "<table>" in markdown
    assert "| 序号 | 参数 | 要求 |" in markdown
    assert not any("价格" in point["original_text"] for point in tables if point["source_category"] == "technical_score")
    score_text = next(point["original_text"] for point in tables if point["source_category"] == "technical_score")
    assert "第二部分客观分" in score_text and "第三部分主观分" in score_text


def test_sources_are_specific_once_and_internal_ids_stay_json_only(tmp_path: Path) -> None:
    data, markdown, _ = load_result(tmp_path)
    points = [point for chapter in data["chapters"] for point in chapter["response_points"]]

    assert all(point["point_id"].startswith("TR-") for point in points)
    assert "TR-" not in markdown
    assert markdown.count("**来源**") == len(data["display_blocks"])
    assert "### 来源" not in markdown
    assert all(ref["file_name"] for point in points for ref in point["source_refs"])
    assert all(ref["clause_path"] or ref["row_or_cell"] for point in points for ref in point["source_refs"])
    assert "页码：未可靠定位" in markdown
    assert not any(ref.get("page_start") for point in points for ref in point["source_refs"])


def test_input_discovery_does_not_read_sibling_reference_directory(tmp_path: Path) -> None:
    input_dir = tmp_path / "input"
    write_hospital_fixture(input_dir)
    reference = tmp_path / "reference_only"
    reference.mkdir()
    Document().save(reference / "不得读取的参考投标文件.docx")

    result = build_technical_outline(input_dir, output_dir=tmp_path / "outputs", enable_page_mapping=False)
    data = json.loads(Path(result["files"]["json"]).read_text(encoding="utf-8"))
    assert all("参考投标文件" not in source["file_name"] for source in data["source_files"])
    assert data["statistics"]["source_file_count"] == 2


def test_display_blocks_preserve_atomic_points_but_group_original_sections(tmp_path: Path) -> None:
    data, markdown, _ = load_result(tmp_path)
    points = {point["point_id"]: point for chapter in data["chapters"] for point in chapter["response_points"]}
    overall = next(chapter for chapter in data["chapters"] if chapter["title"] == "总体技术要求")
    overall_ids = [point["point_id"] for point in overall["response_points"] if point["original_heading"] == "二、技术要求"]
    overall_block = next(block for block in data["display_blocks"] if block["response_point_ids"] == overall_ids)
    service_block = next(block for block in data["display_blocks"] if block["title"] == "（二）服务要求")

    assert len(overall_ids) == 6
    assert len(overall_block["response_point_ids"]) == 6
    assert markdown.count("### 二、技术要求") == 1
    assert markdown.index("（一）投标人须承诺") < markdown.index("（六）具体需求详见")
    assert len(service_block["response_point_ids"]) == 4
    assert all(point_id in points for point_id in service_block["response_point_ids"])
    assert service_block["source_ref"]["block_start"].startswith("paragraph_")
    assert service_block["source_ref"]["block_end"].startswith("paragraph_")


def test_numbered_colon_clause_tracks_continuations_without_merging_atomic_ids(tmp_path: Path) -> None:
    data, markdown, _ = load_result(tmp_path)
    points = [point for chapter in data["chapters"] for point in chapter["response_points"]]
    delivery = next(point for point in points if point["original_text"] == "1. 交货期：")
    continuations = [next(point for point in points if point["point_id"] == point_id) for point_id in delivery["continuation_point_ids"]]

    assert len(continuations) == 2
    assert continuations[0]["original_text"].startswith("货到时间")
    assert continuations[1]["original_text"].startswith("安装完成")
    assert all(point["continuation_of_point_id"] == delivery["point_id"] for point in continuations)
    assert "1. 交货期：\n\n   货到时间" in markdown
    assert "   安装完成" in markdown


def test_marker_notes_and_point_to_point_notes_use_grouped_sources(tmp_path: Path) -> None:
    data, markdown, _ = load_result(tmp_path)
    marker = next(block for block in data["display_blocks"] if block["title"] == "采购清单标记说明")
    form_table = next(block for block in data["display_blocks"] if block["title"] == "技术要求点对点应答表" and block["block_type"] == "table")
    form_notes = next(block for block in data["display_blocks"] if block["title"] == "技术要求点对点应答表填写说明")

    assert len(marker["response_point_ids"]) == 3
    assert form_table["source_ref"]["table_range"]
    assert len(form_table["response_point_ids"]) == 1
    assert len(form_notes["response_point_ids"]) == 2
    assert markdown.count("### 技术要求点对点应答表填写说明") == 1


def test_18_2_and_18_3_are_sibling_clauses_not_headings(tmp_path: Path) -> None:
    _, _, path = load_result(tmp_path)
    blocks = normalize_document_structure(parse_docx(path, {"id": "fixture", "file_ext": "docx", "extract_assets": False})["blocks"])
    by_prefix = {str(block.get("text") or "").split(" ", 1)[0]: block for block in blocks if str(block.get("text") or "").startswith(("18.2 ", "18.3 "))}

    assert by_prefix["18.2"]["is_heading"] is False
    assert by_prefix["18.3"]["is_heading"] is False
    assert by_prefix["18.2"]["node_type"] == by_prefix["18.3"]["node_type"] == "clause"
    assert by_prefix["18.2"]["heading_path"] == by_prefix["18.3"]["heading_path"]
    assert not any("18.2" in item for item in by_prefix["18.3"]["heading_path"])


def test_preview_html_and_display_block_integrity(tmp_path: Path) -> None:
    data, markdown, _ = load_result(tmp_path)
    preview = tmp_path / "outputs" / "input" / "outline" / "technical_response_outline.preview.html"
    html_text = preview.read_text(encoding="utf-8")
    point_index = {point["point_id"]: point for chapter in data["chapters"] for point in chapter["response_points"]}

    assert preview.exists()
    assert '<meta charset="utf-8">' in html_text
    assert "rowspan=" in html_text and "colspan=" in html_text
    assert "class=\"table-wrapper\"" in html_text
    assert "<table>" in markdown
    assert all(point_id in point_index for block in data["display_blocks"] for point_id in block["response_point_ids"])
    for block in data["display_blocks"]:
        files = {point_index[point_id]["source_refs"][0]["file_id"] for point_id in block["response_point_ids"]}
        paths = {tuple(point_index[point_id]["source_refs"][0]["clause_path"]) for point_id in block["response_point_ids"]}
        assert len(files) == len(paths) == 1

    source_keys = [
        (block["source_ref"]["file_id"], tuple(block["source_ref"]["clause_path"]))
        for block in data["display_blocks"]
    ]
    assert not any(source_keys[index] == source_keys[index + 1] == source_keys[index + 2] for index in range(len(source_keys) - 2))
