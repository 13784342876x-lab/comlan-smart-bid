from __future__ import annotations

import copy
import json
from pathlib import Path

from docx import Document

from app.kb_service import KnowledgeBaseService
from app.metadata_store import MetadataStore
from app.simple_special_table_filler import MISSING_VALUE, fill_special_tables, parse_company_performance
from config import METADATA_PATH


PROJECT_NAME = "基于区块链的工业数据管理技术研究与应用服务项目"


def existing_personnel_records() -> list[dict]:
    metadata = json.loads(METADATA_PATH.read_text(encoding="utf-8"))
    records = []
    for record in metadata.get("files", []):
        if isinstance(record.get("resume_parse"), dict) and record["resume_parse"].get("person_name"):
            records.append(copy.deepcopy(record))
        if len(records) == 2:
            break
    assert len(records) == 2
    return records


def configure_temp_storage(monkeypatch, tmp_path: Path) -> None:
    import app.document_parser as document_parser
    import app.file_store as file_store

    storage_dir = tmp_path / "storage"
    monkeypatch.setattr(file_store, "STORAGE_DIR", storage_dir)
    monkeypatch.setattr(file_store, "ORIGINAL_DIR", storage_dir / "original")
    monkeypatch.setattr(document_parser, "PARSED_DIR", storage_dir / "parsed")
    monkeypatch.setattr(document_parser, "ASSETS_DIR", storage_dir / "assets")
    for path in [storage_dir / "original", storage_dir / "parsed", storage_dir / "assets"]:
        path.mkdir(parents=True, exist_ok=True)


def build_test_store(tmp_path: Path, monkeypatch) -> tuple[MetadataStore, list[str]]:
    configure_temp_storage(monkeypatch, tmp_path)
    metadata_path = tmp_path / "metadata.json"
    store = MetadataStore(metadata_path)
    data = store.load()
    data["files"] = existing_personnel_records()
    store.save(data)

    service = KnowledgeBaseService(store)
    fixture = Path(__file__).parent / "fixtures" / "mock_业绩文件_测试.docx"
    record = service.ingest_file(fixture, category="企业业绩", tags=["类似项目业绩", "测试"])
    parsed = service.parse_file(record["id"])
    performance = parse_company_performance(parsed["parsed_text"])
    store.upsert_company_performance(record["id"], performance)
    names = [item["resume_parse"]["person_name"] for item in data["files"][:2]]
    return store, names


def make_special_template(path: Path) -> Path:
    document = Document()
    document.add_paragraph("近年完成的类似项目情况表")
    performance = document.add_table(rows=2, cols=8)
    for index, label in enumerate(
        [
            "项目法人/（服务对象单位名称）",
            "项目名称",
            "项目规模",
            "服务范围",
            "服务期间",
            "项目目前进度",
            "服务人员人数（高峰/平均）",
            "建设单位联系人及电话",
        ]
    ):
        performance.rows[0].cells[index].text = label

    document.add_paragraph("表1 拟委任的主要人员汇总表")
    summary = document.add_table(rows=3, cols=9)
    for index, label in enumerate(["层级", "职务", "姓名", "职称", "证书名称", "级别", "证号", "专业", "备注"]):
        summary.rows[0].cells[index].text = label

    document.add_paragraph("表2主要人员简历表")
    resume = document.add_table(rows=8, cols=9)
    resume.rows[0].cells[0].text = "层级"
    resume.rows[1].cells[0].text = "姓名"
    resume.rows[1].cells[2].text = "性别"
    resume.rows[1].cells[6].text = "年龄"
    resume.rows[2].cells[0].text = "学历"
    resume.rows[2].cells[2].text = "参加工作时间"
    resume.rows[2].cells[6].text = "工作年限"
    resume.rows[3].cells[0].text = "主要资历、经验及承担过的项目"
    for index, label in enumerate(["项目名称", "建设规模", "", "项目起止时间", "", "担负的技术职务", "", "获奖情况", ""]):
        resume.rows[4].cells[index].text = label
    document.save(path)
    return path


def docx_text(path: Path) -> str:
    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def table_text(table) -> str:
    return "\n".join(cell.text for row in table.rows for cell in row.cells)


def resume_tables(document: Document):
    return [
        table
        for table in document.tables
        if "主要资历、经验及承担过的项目" in table_text(table) and "项目名称" in table_text(table)
    ]


def performance_tables(document: Document):
    return [
        table
        for table in document.tables
        if "项目法人" in table_text(table) and "服务人员人数" in table_text(table)
    ]


def summary_tables(document: Document):
    return [
        table
        for table in document.tables
        if all(label in table_text(table) for label in ["层级", "职务", "姓名", "证书名称", "证号", "专业"])
    ]


def has_red_missing_run(document: Document) -> bool:
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        color = run.font.color.rgb
                        if run.text == MISSING_VALUE and color is not None and str(color).upper() == "C00000":
                            return True
    return False


def summary_certificate_names(document: Document) -> list[str]:
    for table in document.tables:
        text = table_text(table)
        if all(label in text for label in ["层级", "职务", "姓名", "证书名称", "证号", "专业"]):
            return [row.cells[4].text for row in table.rows[1:3]]
    return []


def resume_role_values(table) -> list[str]:
    rows = table.rows
    values = []
    for row in rows[5:]:
        if len(row.cells) > 5 and row.cells[5].text.strip():
            values.append(row.cells[5].text.strip())
    return values


def test_fill_special_tables_from_standalone_template(tmp_path: Path, monkeypatch) -> None:
    store, names = build_test_store(tmp_path, monkeypatch)
    template = make_special_template(tmp_path / "业绩人员表.docx")
    output = tmp_path / "业绩人员表_已填充.docx"

    debug = fill_special_tables(template, output, metadata_path=store.path)
    document = Document(str(output))
    text = docx_text(output)
    resumes = resume_tables(document)

    assert output.exists()
    assert (tmp_path / "special_table_fill_debug.json").exists()
    assert PROJECT_NAME in text
    assert all(name in text for name in names)
    assert debug["used_performances"][0]["project_name"] == PROJECT_NAME
    assert text.count("表2主要人员简历表") == 1
    assert len(resumes) == 2
    assert names[0] in table_text(resumes[0])
    assert names[1] in table_text(resumes[1])
    assert "2020年03月-2020年06月" in text
    role_values = [value for table in resumes for value in resume_role_values(table)]
    assert role_values
    assert all(len(value) <= 20 for value in role_values if value != MISSING_VALUE)
    assert "生产环境、办公环境、DMZ" not in "\n".join(role_values)
    assert MISSING_VALUE in text
    assert has_red_missing_run(document)
    assert debug["missing_fields"]
    assert debug["default_used_fields"]
    assert debug["period_normalizations"]
    assert any(item["normalized_period"] == "2020年03月-2020年06月" for item in debug["period_normalizations"])
    assert debug["skipped_long_role_text"] or debug["default_used_fields"]

    certificate_names = summary_certificate_names(document)
    assert certificate_names[0] == MISSING_VALUE
    assert certificate_names[1] != MISSING_VALUE
    assert any(item["person_name"] == names[1] and item.get("certificate_name") for item in debug["certificate_sources"])


def test_fill_special_tables_inside_full_tender_docx(tmp_path: Path, monkeypatch) -> None:
    store, names = build_test_store(tmp_path, monkeypatch)
    template = Path(__file__).parents[1] / "storage" / "original" / "2026" / "招标文件-服务_file_77415c21e6d3.docx"
    output = tmp_path / "招标文件-服务_专项表格已填充.docx"

    fill_special_tables(template, output, metadata_path=store.path)
    document = Document(str(output))
    text = docx_text(output)

    assert output.exists()
    assert text.count("表2主要人员简历表") == 1
    assert len(resume_tables(document)) == 2
    assert PROJECT_NAME in text
    assert all(name in text for name in names)


def test_extract_only_crops_original_document_without_rebuilding_tables(tmp_path: Path, monkeypatch) -> None:
    store, names = build_test_store(tmp_path, monkeypatch)
    template = Path(__file__).parents[1] / "storage" / "original" / "2026" / "招标文件-服务_file_77415c21e6d3.docx"
    output = tmp_path / "专项表格_业绩人员_已填充.docx"

    fill_special_tables(template, output, metadata_path=store.path, extract_only=True)
    original = Document(str(template))
    document = Document(str(output))
    text = docx_text(output)

    assert output.exists()
    assert "近年完成的类似项目情况表" in text
    assert "表1 拟委任的主要人员汇总表" in text
    assert text.count("表2主要人员简历表") == 1
    assert "投标函价格表" not in text
    assert "投标人须知前附表" not in text
    assert len(performance_tables(document)) == 1
    assert len(summary_tables(document)) == 1
    assert len(resume_tables(document)) == 2
    assert PROJECT_NAME in text
    assert all(name in text for name in names)
    original_page_widths = {section.page_width for section in original.sections}
    original_left_margins = {section.left_margin for section in original.sections}
    assert document.sections[0].page_width in original_page_widths
    assert document.sections[0].left_margin in original_left_margins

    for table in document.tables:
        xml = table._tbl.xml
        assert "<w:tblPr" in xml
        assert "<w:tblGrid" in xml
        assert any("<w:tcPr" in cell._tc.xml for row in table.rows for cell in row.cells)
    assert any("gridSpan" in table._tbl.xml or "vMerge" in table._tbl.xml for table in document.tables)


def test_missing_report_lists_personnel_followup_items(tmp_path: Path, monkeypatch) -> None:
    store, names = build_test_store(tmp_path, monkeypatch)
    data = store.load()
    for record in data.get("files", []):
        resume = record.get("resume_parse")
        if isinstance(resume, dict):
            resume["department"] = ""
            resume["source_folder"] = ""
        record.pop("personnel_profile", None)
    store.save(data)

    template = make_special_template(tmp_path / "业绩人员表.docx")
    output = tmp_path / "业绩人员表_已填充.docx"

    debug = fill_special_tables(template, output, metadata_path=store.path)
    document = Document(str(output))
    md_path = output.with_name(f"{output.stem}.missing_report.md")
    json_path = output.with_name(f"{output.stem}.missing_report.json")
    report_text = md_path.read_text(encoding="utf-8")
    report_json = json.loads(json_path.read_text(encoding="utf-8"))

    assert md_path.exists()
    assert json_path.exists()
    assert all(name in report_text for name in names)
    assert report_text.count("### 人员：") == 2
    assert "部门：待确认" in report_text
    for field in ["年龄", "学历", "专业", "参加工作时间"]:
        assert field in report_text
    assert "建议向部门确认" in report_text
    assert has_red_missing_run(document)
    assert debug["missing_fields"]
    assert debug["default_used_fields"]
    assert debug["certificate_sources"]
    assert debug["period_normalizations"]
    assert "missing_report_paths" in debug
    assert any(group["department"] == "待确认" for group in report_json["departments"])
