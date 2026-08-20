from __future__ import annotations

from pathlib import Path

from docx import Document

from app.simple_special_table_filler import fill_special_tables
from app.technical_doc_composer import compose_technical_doc
from tests.test_special_table_filler import PROJECT_NAME, build_test_store, docx_text


def table_text(table) -> str:
    return "\n".join(cell.text for row in table.rows for cell in row.cells)


def test_compose_technical_doc_with_special_tables(tmp_path: Path, monkeypatch) -> None:
    store, names = build_test_store(tmp_path, monkeypatch)
    base = tmp_path / "技术响应正文模板_最小测试.docx"
    tables = tmp_path / "专项表格_业绩人员_已填充.docx"
    output = tmp_path / "技术响应文件_最终版.docx"

    minimal = Document()
    for text in [
        "项目理解",
        "{{PERFORMANCE_TABLES}}",
        "服务方案",
        "{{PROJECT_TEAM_TABLES}}",
        "服务承诺",
    ]:
        minimal.add_paragraph(text)
    minimal.save(base)
    base_doc = Document(str(base))
    assert "项目理解" in "\n".join(paragraph.text for paragraph in base_doc.paragraphs)

    fill_special_tables(
        Path(__file__).parents[1] / "storage" / "original" / "2026" / "招标文件-服务_file_77415c21e6d3.docx",
        tables,
        metadata_path=store.path,
        extract_only=True,
    )
    debug = compose_technical_doc(base, tables, output)

    document = Document(str(output))
    text = docx_text(output)
    table_texts = [table_text(table) for table in document.tables]

    assert output.exists()
    assert "{{PERFORMANCE_TABLES}}" not in text
    assert "{{PROJECT_TEAM_TABLES}}" not in text
    assert "项目理解" in text
    assert "服务方案" in text
    assert "服务承诺" in text
    assert PROJECT_NAME in text
    assert all(name in text for name in names)
    assert any(PROJECT_NAME in value for value in table_texts)
    assert any("拟委任的主要人员汇总表" in text or "证书名称" in value for value in table_texts)
    assert sum(1 for value in table_texts if "主要资历、经验及承担过的项目" in value) == 2
    assert debug["inserted"]["performance_tables"] == 1
    assert debug["inserted"]["project_team_tables"] >= 3
