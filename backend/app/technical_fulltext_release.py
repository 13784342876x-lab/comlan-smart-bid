from __future__ import annotations

from datetime import datetime, timezone
import json
import os
from pathlib import Path
import re
import shutil
import socket
import subprocess
import tempfile
from typing import Any, Callable
from zipfile import ZipFile

import fitz
from docx import Document

from app.technical_doc_source import record_docx_source
from app.technical_fulltext_composer import (
    atomic_write_json,
    atomic_write_markdown,
    load_json_object,
    sha256_file,
)
from app.technical_fulltext_human_review import inspect_technical_fulltext_human_review
from app.technical_fulltext_human_review_finalization import (
    inspect_technical_fulltext_review_finalization,
)
from app.technical_fulltext_hygiene import inspect_technical_fulltext_hygiene
from app.technical_fulltext_release_markdown import render_release_markdown
from app.technical_fulltext_review_package import inspect_technical_fulltext_review_package
from app.technical_review_projection import active_issue_keys, confirmation_groups
from config import BACKEND_ROOT


PANDOC_FALLBACK = Path.home() / "AppData" / "Local" / "Pandoc" / "pandoc.exe"
LIBREOFFICE_PROGRAM = Path(r"C:\Program Files\LibreOffice\program")
REFERENCE_DOCX = (
    BACKEND_ROOT
    / "tools"
    / "pandoc_docx_template"
    / "templates"
    / "template_标题编号-列表第二行顶格.docx"
)
UNO_HELPER = BACKEND_ROOT / "tools" / "refresh_docx_toc.py"
RELEASE_MANIFEST_NAME = "technical_fulltext.release.json"
PLACEHOLDER_PATTERN = re.compile(r"\{\{PH-[^{}]+\}\}")


class TechnicalFulltextReleaseError(RuntimeError):
    pass


class TechnicalFulltextReleaseBlocked(TechnicalFulltextReleaseError):
    def __init__(self, blocking_reasons: list[str]) -> None:
        self.blocking_reasons = blocking_reasons
        super().__init__(f"release blocked: {', '.join(blocking_reasons)}")


class TechnicalFulltextReleaseStale(TechnicalFulltextReleaseError):
    pass


def evaluate_release_gate(
    *,
    cleaned_fulltext_path: str | Path,
    review_package_path: str | Path,
    human_review_path: str | Path,
    finalization_path: str | Path | None,
) -> dict[str, Any]:
    cleaned_path = Path(cleaned_fulltext_path).resolve()
    package_path = Path(review_package_path).resolve()
    human_path = Path(human_review_path).resolve()
    cleaned = load_json_object(cleaned_path, label="cleaned fulltext")
    package = load_json_object(package_path, label="Stage12-A review package")
    human = load_json_object(human_path, label="Stage12-B human review")

    inspect_technical_fulltext_hygiene(cleaned_path)
    inspect_technical_fulltext_review_package(package_path)
    inspect_technical_fulltext_human_review(human_path)

    reasons: list[str] = []
    summary = human["summary"]
    active_keys = active_issue_keys(package, package_path)
    active_pending_issues = [
        item
        for item in human.get("issue_decisions") or []
        if item.get("decision") == "pending"
        and (str(item.get("source_type") or ""), str(item.get("source_id") or ""))
        in active_keys
    ]
    pending_confirmation_groups = [
        item
        for item in confirmation_groups(package, human, package_path)
        if item["status"] == "pending"
    ]
    if active_pending_issues:
        reasons.append("pending_issue_decisions")
    if pending_confirmation_groups:
        reasons.append("pending_fact_confirmations")
    if int(summary["material_verification_counts"]["pending"]) > 0:
        reasons.append("pending_material_verifications")
    if any(not item["accepted"] for item in human.get("risk_acceptances") or []):
        reasons.append("risk_acceptance_required")

    finalization: dict[str, Any] | None = None
    if finalization_path is None:
        reasons.insert(0, "review_not_finalized")
    else:
        final_path = Path(finalization_path).resolve()
        inspected = inspect_technical_fulltext_review_finalization(final_path)
        finalization = load_json_object(final_path, label="Stage12-C finalization")
        identity = finalization["artifact_identity"]
        lineage = finalization["source_lineage"]
        if (
            identity["document_id"] != package["artifact_identity"]["document_id"]
            or identity["output_key"] != cleaned["output_key"]
            or lineage["audited_content_sha256"] != cleaned["cleaned_content_sha256"]
            or human["audited_content_sha256"] != cleaned["cleaned_content_sha256"]
            or package["reviewed_content_sha256"] != cleaned["cleaned_content_sha256"]
            or inspected["source_lineage"]["stage12a_review_package"]["sha256"]
            != sha256_file(package_path)
            or inspected["source_lineage"]["stage12b_human_review"]["sha256"]
            != sha256_file(human_path)
        ):
            reasons.append("lineage_mismatch")
        active_revision_candidates = [
            item
            for item in finalization.get("revision_candidates") or []
            if (str(item.get("source_type") or ""), str(item.get("source_id") or ""))
            in active_keys
        ]
        if active_revision_candidates:
            reasons.append("revision_required")

    reasons = list(dict.fromkeys(reasons))
    return {
        "releaseable": not reasons,
        "blocking_reasons": reasons,
        "output_key": cleaned["output_key"],
        "document_id": cleaned["fulltext"]["document_id"],
        "cleaned_content_sha256": cleaned["cleaned_content_sha256"],
        "finalization_sha256": (
            sha256_file(finalization_path) if finalization is not None else None
        ),
    }


def generate_release_artifacts(
    *,
    cleaned_fulltext_path: str | Path,
    artifact_root: str | Path,
    project_name: str,
    stale_check: Callable[[], None],
) -> dict[str, Any]:
    started = datetime.now(timezone.utc)
    cleaned_path = Path(cleaned_fulltext_path).resolve()
    root = Path(artifact_root).resolve()
    cleaned = load_json_object(cleaned_path, label="cleaned fulltext")
    release_dir = root / "fulltext" / "release"
    release_dir.mkdir(parents=True, exist_ok=True)
    markdown_path = release_dir / "technical_fulltext.release.md"
    docx_path = release_dir / "technical_bid.docx"
    pdf_path = release_dir / "technical_bid.pdf"
    manifest_path = release_dir / RELEASE_MANIFEST_NAME

    markdown = render_release_markdown(cleaned)
    atomic_write_markdown(markdown_path, markdown, required_text="# ")
    pandoc = _pandoc_executable()
    pandoc_started = datetime.now(timezone.utc)
    subprocess.run(
        [
            str(pandoc),
            str(markdown_path),
            "--from=gfm",
            "--to=docx",
            "--toc",
            "--toc-depth=3",
            "-M",
            "lang=zh-CN",
            "-M",
            "toc-title=目录",
            "-M",
            f"title={project_name}_投标文件",
            f"--reference-doc={REFERENCE_DOCX}",
            f"--output={docx_path}",
        ],
        check=True,
        capture_output=True,
        text=True,
        encoding="utf-8",
    )
    pandoc_finished = datetime.now(timezone.utc)
    record_docx_source(
        markdown_path,
        docx_path,
        converter=_version_line(pandoc, "pandoc"),
        reference_doc=REFERENCE_DOCX,
    )

    office_started = datetime.now(timezone.utc)
    _refresh_toc_and_export_pdf(docx_path, pdf_path)
    office_finished = datetime.now(timezone.utc)
    docx_validation = _validate_docx(docx_path, cleaned, markdown)
    pdf_validation = _validate_pdf(pdf_path, docx_validation)
    validation_finished = datetime.now(timezone.utc)

    stale_check()
    generated_at = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")
    manifest = {
        "schema_version": "technical_fulltext_release_v1",
        "status": "published",
        "output_key": cleaned["output_key"],
        "document_id": cleaned["fulltext"]["document_id"],
        "source": {
            "cleaned_fulltext_path": str(cleaned_path.relative_to(root)),
            "cleaned_fulltext_sha256": sha256_file(cleaned_path),
            "cleaned_content_sha256": cleaned["cleaned_content_sha256"],
        },
        "markdown": {
            "path": str(markdown_path.relative_to(root)),
            "sha256": sha256_file(markdown_path),
        },
        "reference_docx": {
            "path": str(REFERENCE_DOCX.resolve()),
            "sha256": sha256_file(REFERENCE_DOCX),
        },
        "pandoc_version": _version_line(pandoc, "pandoc"),
        "libreoffice_version": _version_line(
            LIBREOFFICE_PROGRAM / "soffice.com", "LibreOffice"
        ),
        "docx": {
            "path": str(docx_path.relative_to(root)),
            "sha256": sha256_file(docx_path),
            "download_filename": f"{_safe_filename(project_name)}_技术标正文.docx",
            "validation": docx_validation,
        },
        "pdf": {
            "path": str(pdf_path.relative_to(root)),
            "sha256": sha256_file(pdf_path),
            "download_filename": f"{_safe_filename(project_name)}_技术标正文.pdf",
            "validation": pdf_validation,
        },
        "timing_seconds": {
            "pandoc_docx": round((pandoc_finished - pandoc_started).total_seconds(), 3),
            "toc_refresh_and_pdf": round((office_finished - office_started).total_seconds(), 3),
            "validation": round((validation_finished - office_finished).total_seconds(), 3),
            "total": round((datetime.now(timezone.utc) - started).total_seconds(), 3),
        },
        "generated_at": generated_at,
    }
    atomic_write_json(
        manifest_path,
        manifest,
        required_fields=(
            "schema_version",
            "status",
            "output_key",
            "source",
            "markdown",
            "reference_docx",
            "docx",
            "pdf",
            "generated_at",
        ),
    )
    return {**manifest, "manifest_path": str(manifest_path)}


def inspect_release_manifest(
    *, artifact_root: str | Path, expected_output_key: str, expected_cleaned_sha256: str
) -> dict[str, Any] | None:
    root = Path(artifact_root).resolve()
    path = root / "fulltext" / "release" / RELEASE_MANIFEST_NAME
    if not path.is_file():
        return None
    manifest = load_json_object(path, label="technical fulltext release")
    if (
        manifest.get("status") != "published"
        or manifest.get("output_key") != expected_output_key
        or (manifest.get("source") or {}).get("cleaned_content_sha256")
        != expected_cleaned_sha256
    ):
        raise TechnicalFulltextReleaseError("release manifest does not bind current artifact")
    for key in ("markdown", "docx", "pdf"):
        record = manifest[key]
        target = (root / record["path"]).resolve()
        target.relative_to(root)
        if not target.is_file() or sha256_file(target) != record["sha256"]:
            raise TechnicalFulltextReleaseError(f"release {key} artifact is invalid")
    return {**manifest, "manifest_path": str(path)}


def _pandoc_executable() -> Path:
    discovered = shutil.which("pandoc")
    target = Path(discovered) if discovered else PANDOC_FALLBACK
    if not target.is_file():
        raise TechnicalFulltextReleaseError("Pandoc executable is unavailable")
    if not REFERENCE_DOCX.is_file():
        raise TechnicalFulltextReleaseError("fixed reference DOCX is unavailable")
    return target.resolve()


def _refresh_toc_and_export_pdf(docx_path: Path, pdf_path: Path) -> None:
    soffice = LIBREOFFICE_PROGRAM / "soffice.com"
    python = LIBREOFFICE_PROGRAM / "python.exe"
    if not all(path.is_file() for path in (soffice, python, UNO_HELPER)):
        raise TechnicalFulltextReleaseError("LibreOffice UNO runtime is unavailable")
    port = _available_office_port()
    with tempfile.TemporaryDirectory(
        prefix="comlan-release-lo-", ignore_cleanup_errors=True
    ) as profile:
        profile_url = Path(profile).resolve().as_uri()
        command = (
            f"& '{soffice}' --headless --invisible --nologo --nodefault "
            f"--norestore --nofirststartwizard "
            f"'-env:UserInstallation={profile_url}' "
            f"'--accept=socket,host=127.0.0.1,port={port};urp;'"
        )
        process = subprocess.Popen(
            [
                r"C:\Windows\System32\WindowsPowerShell\v1.0\powershell.exe",
                "-NoProfile",
                "-Command",
                command,
            ],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
        try:
            subprocess.run(
                [
                    str(python),
                    str(UNO_HELPER),
                    "--port",
                    str(port),
                    "--docx",
                    str(docx_path),
                    "--pdf",
                    str(pdf_path),
                ],
                check=True,
                timeout=180,
            )
        finally:
            if process.poll() is None:
                process.terminate()
                try:
                    process.wait(timeout=10)
                except subprocess.TimeoutExpired:
                    process.kill()
                    process.wait(timeout=10)


def _validate_docx(
    docx_path: Path, cleaned: dict[str, Any], markdown: str
) -> dict[str, Any]:
    if not docx_path.is_file() or docx_path.stat().st_size == 0:
        raise TechnicalFulltextReleaseError("DOCX was not generated")
    document = Document(docx_path)
    chapters = sorted(cleaned["fulltext"]["chapters"], key=lambda item: item["order"])
    section_titles = [
        section["heading"]
        for chapter in chapters
        for section in sorted(
            chapter["sections"],
            key=lambda item: next(
                outline["document_order"]
                for outline in chapter["section_outline"]
                if outline["section_id"] == item["section_id"]
            ),
        )
    ]
    paragraph_text = [paragraph.text for paragraph in document.paragraphs]
    missing_titles = [
        title
        for title in [*(chapter["chapter_title"] for chapter in chapters), *section_titles]
        if title not in paragraph_text
    ]
    if missing_titles:
        raise TechnicalFulltextReleaseError(
            f"DOCX is missing canonical headings: {missing_titles[:3]}"
        )
    placeholders = sorted(set(PLACEHOLDER_PATTERN.findall(markdown)))
    document_text = "\n".join(paragraph_text)
    for table in document.tables:
        document_text += "\n" + "\n".join(
            cell.text for row in table.rows for cell in row.cells
        )
    missing_placeholders = [value for value in placeholders if value not in document_text]
    if missing_placeholders:
        raise TechnicalFulltextReleaseError(
            f"DOCX is missing placeholders: {missing_placeholders[:3]}"
        )
    with ZipFile(docx_path) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        field_xml = "\n".join(
            archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if name.startswith("word/") and name.endswith(".xml")
        )
    toc_entries = document_xml.count("<w:hyperlink")
    if " TOC " not in document_xml or toc_entries < len(chapters):
        raise TechnicalFulltextReleaseError("DOCX TOC has no cached entries")
    if " PAGE " not in field_xml:
        raise TechnicalFulltextReleaseError("DOCX page-number field is missing")
    return {
        "chapter_count": len(chapters),
        "section_count": len(section_titles),
        "table_count": len(document.tables),
        "placeholder_count": len(placeholders),
        "toc_entry_count": toc_entries,
        "page_field_present": True,
        "size_bytes": docx_path.stat().st_size,
    }


def _validate_pdf(pdf_path: Path, docx_validation: dict[str, Any]) -> dict[str, Any]:
    if not pdf_path.is_file() or pdf_path.stat().st_size == 0:
        raise TechnicalFulltextReleaseError("PDF was not generated")
    document = fitz.open(pdf_path)
    try:
        if document.page_count == 0:
            raise TechnicalFulltextReleaseError("PDF contains no pages")
        text = "\n".join(page.get_text() for page in document)
        if len(text.strip()) < 500:
            raise TechnicalFulltextReleaseError("PDF contains insufficient body text")
        toc = document.get_toc()
        expected = docx_validation["chapter_count"] + docx_validation["section_count"]
        if len(toc) < expected or not all(int(item[2]) > 0 for item in toc):
            raise TechnicalFulltextReleaseError("PDF TOC entries or page references are missing")
        first_pages = "\n".join(
            document[index].get_text() for index in range(min(4, document.page_count))
        )
        if not re.search(r"\.{5,}\s*\d+", first_pages):
            raise TechnicalFulltextReleaseError("PDF TOC page references are not rendered")
        footer_numbers = sum(
            1
            for page in document
            if any(
                word[4].isdigit() and float(word[1]) > page.rect.height * 0.9
                for word in page.get_text("words")
            )
        )
        if footer_numbers == 0:
            raise TechnicalFulltextReleaseError("PDF footer page numbers are missing")
        return {
            "page_count": document.page_count,
            "outline_entry_count": len(toc),
            "text_character_count": len(text),
            "footer_page_number_pages": footer_numbers,
            "size_bytes": pdf_path.stat().st_size,
        }
    finally:
        document.close()


def _version_line(executable: Path, label: str) -> str:
    result = subprocess.run(
        [str(executable), "--version"],
        check=True,
        capture_output=True,
        text=True,
        encoding="utf-8",
    )
    line = result.stdout.splitlines()[0].strip() if result.stdout else ""
    if not line:
        raise TechnicalFulltextReleaseError(f"{label} version is unavailable")
    return line


def _safe_filename(value: str) -> str:
    safe = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", value).strip(" .")
    return safe[:120] or "项目"


def _available_office_port() -> int:
    for port in range(2100, 3100):
        with socket.socket() as probe:
            try:
                probe.bind(("127.0.0.1", port))
            except OSError:
                continue
        return port
    raise TechnicalFulltextReleaseError("no local LibreOffice UNO port is available")
