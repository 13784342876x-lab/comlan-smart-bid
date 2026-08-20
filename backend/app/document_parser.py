import shutil
import subprocess
import uuid
import zipfile
from pathlib import Path
from typing import Any

from app.file_store import resolve_storage_key
from config import ASSETS_DIR, PARSED_DIR, ensure_directories

IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "webp"}
TEXT_EXTENSIONS = {"txt", "md"}
WORD_EXTENSIONS = {"doc", "docx"}
EXCEL_EXTENSIONS = {"xls", "xlsx"}
PPT_EXTENSIONS = {"ppt", "pptx"}

SOFFICE_CANDIDATES = [
    Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
    Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"),
]


def parse_document(record: dict[str, Any]) -> dict[str, Any]:
    ensure_directories()
    ext = record["file_ext"].lower()
    source = resolve_storage_key(record["storage_key"])
    result = empty_result(record)

    try:
        if ext == "docx":
            result.update(parse_docx(source, record))
        elif ext == "doc":
            result.update(parse_via_libreoffice(source, record, "docx", parse_docx))
        elif ext == "pdf":
            result.update(parse_pdf(source, record))
        elif ext == "xlsx":
            result.update(parse_xlsx(source, record))
        elif ext == "xls":
            result.update(parse_via_libreoffice(source, record, "xlsx", parse_xlsx))
        elif ext == "pptx":
            result.update(parse_pptx(source, record))
        elif ext == "ppt":
            result.update(parse_via_libreoffice(source, record, "pptx", parse_pptx))
        elif ext in TEXT_EXTENSIONS:
            result.update(parse_text_file(source))
        elif ext in IMAGE_EXTENSIONS:
            result.update(
                {
                    "ok": False,
                    "message": "图片文件需要 OCR 模块，当前暂未启用",
                    "warnings": ["图片文件需要 OCR 模块，当前暂未启用"],
                }
            )
        else:
            result.update(
                {
                    "ok": False,
                    "message": f"暂不支持解析 .{ext} 文件",
                    "warnings": [f"unsupported parser for .{ext}"],
                }
            )
    except Exception as exc:
        result.update(
            {
                "ok": False,
                "message": f"文件解析失败: {exc}",
                "warnings": [f"parse failed: {exc}"],
            }
        )

    if result.get("parsed_text"):
        (PARSED_DIR / f"{record['id']}.txt").write_text(result["parsed_text"], encoding="utf-8")
    return result


def empty_result(record: dict[str, Any]) -> dict[str, Any]:
    return {
        "ok": False,
        "file_id": record["id"],
        "file_type": record["file_ext"].lower(),
        "parsed_text": "",
        "sections": [],
        "tables": [],
        "blocks": [],
        "assets": [],
        "warnings": [],
        "message": "",
    }


def parse_text_file(path: Path) -> dict[str, Any]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    blocks = [
        {
            "block_index": index,
            "block_type": "paragraph",
            "source_location": f"line_{index}",
            "text": line.strip(),
        }
        for index, line in enumerate(text.splitlines(), start=1)
        if line.strip()
    ]
    warnings = [] if text.strip() else ["文本文件为空"]
    return {
        "ok": bool(text.strip()),
        "parsed_text": text,
        "sections": [],
        "tables": [],
        "blocks": blocks,
        "assets": [],
        "warnings": warnings,
        "message": "解析完成" if text.strip() else "文本文件为空",
    }


def parse_docx(path: Path, record: dict[str, Any]) -> dict[str, Any]:
    try:
        from docx import Document
    except ImportError as exc:
        return dependency_missing("python-docx", exc)

    document = Document(str(path))
    sections: list[dict[str, Any]] = []
    tables: list[dict[str, Any]] = []
    blocks: list[dict[str, Any]] = []
    text_parts: list[str] = []
    paragraph_index = 0
    table_index = 0
    for block_index, (block_type, block) in enumerate(iter_docx_body_blocks(document), start=1):
        if block_type == "paragraph":
            paragraph_index += 1
            text = block.text.strip()
            if not text:
                continue
            structure = docx_paragraph_structure(block)
            section = {"source_location": f"paragraph_{paragraph_index}", "text": text, **structure}
            sections.append(section)
            blocks.append(
                {
                    "block_index": block_index,
                    "block_type": "paragraph",
                    "source_location": section["source_location"],
                    "text": text,
                    **structure,
                }
            )
            text_parts.append(text)
            continue
        table_index += 1
        table_data = extract_docx_table_grid(block, table_index)
        tables.append(table_data)
        blocks.append(
            {
                "block_index": block_index,
                "block_type": "table",
                "table_index": table_index,
                "source_location": f"table_{table_index}",
                **table_data,
            }
        )
        text_parts.append(f"=== Table {table_index} ===")
        text_parts.extend("\t".join(row) for row in table_data["rows"])

    assets = extract_docx_images(path, record) if record.get("extract_assets", True) else []
    parsed_text = "\n".join(text_parts).strip()
    warnings = [] if parsed_text else ["DOCX 未提取到文本内容"]
    return {
        "ok": bool(parsed_text),
        "parsed_text": parsed_text,
        "sections": sections,
        "tables": tables,
        "blocks": blocks,
        "assets": assets,
        "warnings": warnings,
        "message": "解析完成" if parsed_text else "DOCX 未提取到文本内容",
    }


def docx_paragraph_structure(paragraph) -> dict[str, Any]:
    from docx.oxml.ns import qn

    p_pr = paragraph._p.pPr
    outline_level = None
    numbering_level = None
    numbering_id = None
    if p_pr is not None:
        outline = p_pr.find(qn("w:outlineLvl"))
        if outline is not None and outline.get(qn("w:val")) is not None:
            outline_level = int(outline.get(qn("w:val")))
        num_pr = p_pr.find(qn("w:numPr"))
        if num_pr is not None:
            level = num_pr.find(qn("w:ilvl"))
            num_id = num_pr.find(qn("w:numId"))
            if level is not None and level.get(qn("w:val")) is not None:
                numbering_level = int(level.get(qn("w:val")))
            if num_id is not None and num_id.get(qn("w:val")) is not None:
                numbering_id = int(num_id.get(qn("w:val")))

    sizes = [run.font.size.pt for run in paragraph.runs if run.font.size is not None]
    direct_bold = [run.bold for run in paragraph.runs if run.text.strip()]
    style_font = getattr(getattr(paragraph, "style", None), "font", None)
    style_size = style_font.size.pt if style_font is not None and style_font.size is not None else None
    style_bold = style_font.bold if style_font is not None else None
    fmt = paragraph.paragraph_format
    return {
        "original_style": str(getattr(getattr(paragraph, "style", None), "name", "") or ""),
        "outline_level": outline_level,
        "numbering_level": numbering_level,
        "numbering_id": numbering_id,
        "font_size_pt": max(sizes) if sizes else style_size,
        "is_bold": bool(any(value is True for value in direct_bold) or style_bold is True),
        "alignment": int(paragraph.alignment) if paragraph.alignment is not None else None,
        "space_before_pt": fmt.space_before.pt if fmt.space_before is not None else None,
        "space_after_pt": fmt.space_after.pt if fmt.space_after is not None else None,
        "keep_with_next": fmt.keep_with_next,
    }


def extract_docx_table_grid(table, table_index: int) -> dict[str, Any]:
    from docx.oxml.ns import qn
    from docx.table import _Cell

    column_count = len(table._tbl.tblGrid.gridCol_lst) if table._tbl.tblGrid is not None else len(table.columns)
    structured_rows: list[dict[str, Any]] = []
    logical_rows: list[list[str]] = []
    active_vertical: dict[int, dict[str, Any]] = {}
    has_merged_cells = False

    for row_number, tr in enumerate(table._tbl.tr_lst, start=1):
        cells: list[dict[str, Any]] = []
        logical = [""] * max(1, column_count)
        column = 1
        extended_anchors: set[int] = set()
        for tc in tr.tc_lst:
            tc_pr = tc.tcPr
            grid_span = tc_pr.find(qn("w:gridSpan")) if tc_pr is not None else None
            v_merge = tc_pr.find(qn("w:vMerge")) if tc_pr is not None else None
            colspan = int(grid_span.get(qn("w:val"))) if grid_span is not None else 1
            merge_value = None
            if v_merge is not None:
                merge_value = v_merge.get(qn("w:val")) or "continue"
            if colspan > 1 or v_merge is not None:
                has_merged_cells = True

            covered_columns = range(column, column + colspan)
            if merge_value == "continue":
                anchor = active_vertical.get(column)
                if anchor is not None and id(anchor) not in extended_anchors:
                    anchor["rowspan"] += 1
                    extended_anchors.add(id(anchor))
                for covered in covered_columns:
                    if anchor is not None:
                        active_vertical[covered] = anchor
                column += colspan
                continue

            text = _Cell(tc, table).text.strip()
            cell = {
                "text": text,
                "row": row_number,
                "column": column,
                "rowspan": 1,
                "colspan": colspan,
                "is_merge_anchor": True,
            }
            cells.append(cell)
            if column - 1 >= len(logical):
                logical.extend([""] * (column - len(logical)))
            logical[column - 1] = text
            for covered in covered_columns:
                if merge_value == "restart":
                    active_vertical[covered] = cell
                else:
                    active_vertical.pop(covered, None)
            column += colspan

        structured_rows.append({"row": row_number, "cells": cells})
        logical_rows.append(logical)

    return {
        "table_index": table_index,
        "rows": logical_rows,
        "structured_rows": structured_rows,
        "column_count": column_count,
        "has_merged_cells": has_merged_cells,
        "source_location": f"table_{table_index}",
    }


def iter_docx_body_blocks(document):
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield "paragraph", Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield "table", Table(child, document)


def extract_docx_images(path: Path, record: dict[str, Any]) -> list[dict[str, Any]]:
    assets: list[dict[str, Any]] = []
    with zipfile.ZipFile(path) as archive:
        media_names = [name for name in archive.namelist() if name.startswith("word/media/")]
        for index, media_name in enumerate(media_names, start=1):
            suffix = Path(media_name).suffix.lower() or ".bin"
            asset_id = f"asset_{uuid.uuid4().hex[:12]}"
            asset_name = f"{record['id']}_{asset_id}{suffix}"
            target = ASSETS_DIR / asset_name
            with archive.open(media_name) as source, target.open("wb") as output:
                shutil.copyfileobj(source, output)
            assets.append(
                {
                    "id": asset_id,
                    "file_id": record["id"],
                    "name": Path(media_name).name,
                    "asset_type": "image",
                    "source_page": f"docx_media_{index}",
                    "source_location": f"word/media/{Path(media_name).name}",
                    "asset_key": f"assets/{asset_name}",
                    "preview_url": f"assets/{asset_name}",
                    "ocr_text": "",
                    "recommended_usage": "源文件图片资产，可在后续标书生成中按需引用",
                    "insertable": True,
                    "insert_mode": "original_ratio",
                }
            )
    return assets


def parse_pdf(path: Path, record: dict[str, Any]) -> dict[str, Any]:
    try:
        import fitz
    except ImportError as exc:
        return dependency_missing("PyMuPDF", exc)

    pages: list[dict[str, Any]] = []
    text_parts: list[str] = []
    with fitz.open(path) as document:
        for page_index, page in enumerate(document, start=1):
            text = page.get_text("text").strip()
            pages.append({"page_number": page_index, "text": text})
            text_parts.append(f"=== Page {page_index} ===")
            text_parts.append(text)

    parsed_text = "\n".join(text_parts).strip()
    plain_text_len = sum(len(page["text"]) for page in pages)
    warnings = []
    if plain_text_len < 80:
        warnings.append("疑似扫描件 PDF，后续需要 OCR")
    return {
        "ok": bool(parsed_text),
        "parsed_text": parsed_text,
        "sections": pages,
        "tables": [],
        "blocks": [
            {
                "block_index": page["page_number"],
                "block_type": "page",
                "page_number": page["page_number"],
                "source_location": f"page_{page['page_number']}",
                "text": page["text"],
            }
            for page in pages
        ],
        "assets": [],
        "warnings": warnings,
        "message": "解析完成" if parsed_text else "PDF 未提取到文本内容",
    }


def parse_xlsx(path: Path, record: dict[str, Any]) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        return dependency_missing("openpyxl", exc)

    workbook = load_workbook(path, data_only=True)
    tables: list[dict[str, Any]] = []
    text_parts: list[str] = []
    blocks: list[dict[str, Any]] = []
    for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
        rows: list[dict[str, Any]] = []
        text_parts.append(f"=== Sheet: {sheet.title} ===")
        for row in sheet.iter_rows():
            values = [cell.value for cell in row]
            if all(value is None for value in values):
                continue
            cells = []
            line_values = []
            for cell in row:
                value = "" if cell.value is None else str(cell.value)
                cells.append({"row": cell.row, "col": cell.column, "value": value})
                line_values.append(value)
            rows.append({"row": row[0].row, "cells": cells})
            text_parts.append("\t".join(line_values))
        table_data = {
            "sheet_name": sheet.title,
            "merged_cells": [str(item) for item in sheet.merged_cells.ranges],
            "rows": rows,
        }
        tables.append(table_data)
        blocks.append(
            {
                "block_index": sheet_index,
                "block_type": "table",
                "sheet_name": sheet.title,
                "source_location": f"sheet_{sheet.title}",
                **table_data,
            }
        )

    parsed_text = "\n".join(text_parts).strip()
    warnings = [] if parsed_text else ["Excel 未提取到文本内容"]
    return {
        "ok": bool(parsed_text),
        "parsed_text": parsed_text,
        "sections": [],
        "tables": tables,
        "blocks": blocks,
        "assets": [],
        "warnings": warnings,
        "message": "解析完成" if parsed_text else "Excel 未提取到文本内容",
    }


def parse_pptx(path: Path, record: dict[str, Any]) -> dict[str, Any]:
    try:
        from pptx import Presentation
    except ImportError as exc:
        return dependency_missing("python-pptx", exc)

    presentation = Presentation(str(path))
    slides: list[dict[str, Any]] = []
    tables: list[dict[str, Any]] = []
    text_parts: list[str] = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        slide_lines: list[str] = []
        text_parts.append(f"=== Slide {slide_index} ===")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_lines.append(shape.text.strip())
            if getattr(shape, "has_table", False):
                rows = []
                for row in shape.table.rows:
                    row_values = [cell.text.strip() for cell in row.cells]
                    rows.append(row_values)
                    slide_lines.append("\t".join(row_values))
                tables.append({"slide_number": slide_index, "rows": rows})
        slides.append({"slide_number": slide_index, "text": "\n".join(slide_lines)})
        text_parts.extend(slide_lines)

    parsed_text = "\n".join(text_parts).strip()
    warnings = [] if parsed_text else ["PPTX 未提取到文本内容"]
    return {
        "ok": bool(parsed_text),
        "parsed_text": parsed_text,
        "sections": slides,
        "tables": tables,
        "blocks": [
            {
                "block_index": slide["slide_number"],
                "block_type": "slide",
                "slide_number": slide["slide_number"],
                "source_location": f"slide_{slide['slide_number']}",
                "text": slide["text"],
            }
            for slide in slides
        ],
        "assets": [],
        "warnings": warnings,
        "message": "解析完成" if parsed_text else "PPTX 未提取到文本内容",
    }


def parse_via_libreoffice(path: Path, record: dict[str, Any], target_ext: str, parser) -> dict[str, Any]:
    soffice = find_soffice()
    if not soffice:
        return {
            "ok": False,
            "parsed_text": "",
            "sections": [],
            "tables": [],
            "assets": [],
            "warnings": [
                f"当前解析 .{record['file_ext']} 需要安装 LibreOffice，或请先另存为 .{target_ext}"
            ],
            "message": f"当前解析 .{record['file_ext']} 需要安装 LibreOffice，或请先另存为 .{target_ext}",
        }

    output_dir = PARSED_DIR / "converted"
    output_dir.mkdir(parents=True, exist_ok=True)
    command = [
        str(soffice),
        "--headless",
        "--convert-to",
        target_ext,
        "--outdir",
        str(output_dir),
        str(path),
    ]
    completed = subprocess.run(command, capture_output=True, text=True, timeout=120)
    if completed.returncode != 0:
        return {
            "ok": False,
            "parsed_text": "",
            "sections": [],
            "tables": [],
            "assets": [],
            "warnings": [completed.stderr.strip() or completed.stdout.strip() or "LibreOffice 转换失败"],
            "message": f"LibreOffice 转换 .{record['file_ext']} 失败",
        }
    converted = output_dir / f"{path.stem}.{target_ext}"
    if not converted.exists():
        candidates = sorted(output_dir.glob(f"*.{target_ext}"), key=lambda item: item.stat().st_mtime, reverse=True)
        converted = candidates[0] if candidates else converted
    if not converted.exists():
        return {
            "ok": False,
            "parsed_text": "",
            "sections": [],
            "tables": [],
            "assets": [],
            "warnings": ["LibreOffice 转换后未找到输出文件"],
            "message": f"LibreOffice 转换 .{record['file_ext']} 失败",
        }
    converted_record = {**record, "file_ext": target_ext}
    return parser(converted, converted_record)


def find_soffice() -> Path | None:
    for candidate in SOFFICE_CANDIDATES:
        if candidate.exists():
            return candidate
    found = shutil.which("soffice") or shutil.which("libreoffice")
    return Path(found) if found else None


def dependency_missing(package: str, exc: ImportError) -> dict[str, Any]:
    return {
        "ok": False,
        "parsed_text": "",
        "sections": [],
        "tables": [],
        "assets": [],
        "warnings": [f"缺少依赖 {package}: {exc}"],
        "message": f"缺少依赖 {package}，请先安装 requirements.txt",
    }
