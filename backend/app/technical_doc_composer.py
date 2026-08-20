from __future__ import annotations

import copy
import json
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

from app.simple_special_table_filler import (
    is_performance_table,
    is_resume_table,
    is_summary_table,
    table_contexts,
)


PERFORMANCE_PLACEHOLDER = "{{PERFORMANCE_TABLES}}"
PROJECT_TEAM_PLACEHOLDER = "{{PROJECT_TEAM_TABLES}}"


def compose_technical_doc(
    base_path: str | Path,
    tables_path: str | Path,
    output_path: str | Path,
    debug_path: str | Path | None = None,
) -> dict[str, Any]:
    base_path = Path(base_path)
    tables_path = Path(tables_path)
    output_path = Path(output_path)
    debug_path = Path(debug_path) if debug_path else output_path.parent / "technical_doc_compose_debug.json"

    base = Document(str(base_path))
    tables_doc = Document(str(tables_path))
    groups = collect_table_groups(tables_doc)

    debug: dict[str, Any] = {
        "base": str(base_path),
        "tables": str(tables_path),
        "output": str(output_path),
        "inserted": {
            "performance_tables": len(groups["performance"]),
            "project_team_tables": len(groups["project_team"]),
        },
        "missing_placeholders": [],
    }

    if not replace_placeholder_with_tables(base, PERFORMANCE_PLACEHOLDER, groups["performance"]):
        debug["missing_placeholders"].append(PERFORMANCE_PLACEHOLDER)
    if not replace_placeholder_with_tables(base, PROJECT_TEAM_PLACEHOLDER, groups["project_team"]):
        debug["missing_placeholders"].append(PROJECT_TEAM_PLACEHOLDER)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    base.save(str(output_path))
    debug["debug_path"] = str(debug_path)
    debug_path.parent.mkdir(parents=True, exist_ok=True)
    debug_path.write_text(json.dumps(debug, ensure_ascii=False, indent=2), encoding="utf-8")
    return debug


def collect_table_groups(document: Document) -> dict[str, list[Any]]:
    groups = {"performance": [], "project_team": []}
    for item in table_contexts(document):
        table = item["table"]
        text = item["text"]
        if is_performance_table(text):
            groups["performance"].append(table)
        elif is_summary_table(text) or is_resume_table(text):
            groups["project_team"].append(table)
    return groups


def replace_placeholder_with_tables(document: Document, placeholder: str, tables: list[Any]) -> bool:
    for paragraph in document.paragraphs:
        if placeholder not in paragraph.text:
            continue
        clear_paragraph_text(paragraph)
        anchor = paragraph._p
        for index, table in enumerate(tables):
            if index:
                blank = OxmlElement("w:p")
                anchor.addnext(blank)
                anchor = blank
            new_table = copy.deepcopy(table._tbl)
            anchor.addnext(new_table)
            anchor = new_table
        return True
    return False


def clear_paragraph_text(paragraph: Paragraph) -> None:
    for run in paragraph.runs:
        run.text = run.text.replace(PERFORMANCE_PLACEHOLDER, "").replace(PROJECT_TEAM_PLACEHOLDER, "")
