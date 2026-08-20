from __future__ import annotations

import copy
import json
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.shared import Pt, RGBColor
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from app.metadata_store import MetadataStore
from config import METADATA_PATH


MISSING_VALUE = "待补充"
MISSING_COLOR = RGBColor(0xC0, 0x00, 0x00)
FILL_FONT_SIZE = Pt(9)


PERFORMANCE_FIELDS = {
    "owner_unit": ["项目法人/服务对象单位名称", "项目法人", "服务对象单位名称", "服务对象单位"],
    "project_name": ["项目名称"],
    "project_scale": ["项目规模", "建设规模"],
    "service_scope": ["服务范围"],
    "service_period": ["服务期间", "项目起止时间"],
    "progress": ["项目目前进度", "目前进度"],
    "service_staff_count": ["服务人员人数", "高峰/平均"],
    "contact": ["建设单位联系人及电话", "服务对象联系信息", "联系人及电话"],
}

PERFORMANCE_TABLE_FIELDS = [
    "owner_unit",
    "project_name",
    "project_scale",
    "service_scope",
    "service_period",
    "progress",
    "service_staff_count",
    "contact",
]

SUMMARY_FIELDS = [
    "level",
    "job",
    "name",
    "title",
    "certificate_name",
    "certificate_level",
    "certificate_no",
    "certificate_major",
    "remark",
]

SUMMARY_LABELS = {
    "level": ["层级"],
    "job": ["职务"],
    "name": ["姓名"],
    "title": ["职称"],
    "certificate_name": ["证书名称"],
    "certificate_level": ["级别"],
    "certificate_no": ["证号", "证书编号"],
    "certificate_major": ["专业"],
    "remark": ["备注"],
}

RESUME_LABELS = {
    "level": ["层级"],
    "name": ["姓名"],
    "gender": ["性别"],
    "age": ["年龄"],
    "education": ["学历"],
    "start_work_date": ["参加工作时间"],
    "work_years": ["工作年限"],
}

PROJECT_LABELS = {
    "name": ["项目名称"],
    "scale": ["建设规模"],
    "date": ["项目起止时间", "项目日期", "项目时间"],
    "role": ["担负的技术职务"],
    "award": ["获奖情况"],
}

FIELD_REPORT_LABELS = {
    "certificate_no": "证号",
    "certificate_major": "专业",
    "age": "年龄",
    "education": "学历",
    "start_work_date": "参加工作时间",
    "work_years": "工作年限",
    "award": "获奖情况",
    "role": "担负的技术职务",
    "title": "职称",
    "certificate_name": "证书名称",
    "certificate_level": "证书级别",
    "job": "职务",
    "date": "项目起止时间",
    "scale": "建设规模",
}

FIELD_USAGE = {
    "certificate_no": "用于主要人员汇总表的执业或职业资格证明",
    "certificate_major": "用于主要人员汇总表的专业字段",
    "age": "用于主要人员简历表",
    "education": "用于主要人员简历表",
    "start_work_date": "用于主要人员简历表",
    "work_years": "用于主要人员简历表",
    "award": "用于主要人员简历表中的项目经历",
    "role": "用于主要人员简历表中的项目经历",
    "title": "用于主要人员汇总表",
    "certificate_name": "用于主要人员汇总表",
    "certificate_level": "用于主要人员汇总表",
    "job": "用于主要人员汇总表",
    "date": "用于主要人员简历表中的项目经历",
    "scale": "用于主要人员简历表中的项目经历",
}

FIELD_QUESTIONS = {
    "certificate_no": "请提供对应证书编号。",
    "certificate_major": "请确认证书专业或人员所学专业。",
    "age": "请提供出生年月或年龄。",
    "education": "请提供最高学历。",
    "start_work_date": "请提供参加工作时间。",
    "work_years": "请确认工作年限。",
    "award": "请确认相关项目是否有获奖情况；如无，请明确回复无。",
    "role": "请确认该项目中担负的技术职务。",
    "title": "请提供职称信息。",
    "certificate_name": "请提供可用于本次投标的证书名称。",
    "certificate_level": "请提供证书级别。",
    "job": "请确认本次投标职务。",
    "date": "请确认项目起止时间。",
    "scale": "请补充项目建设规模。",
}

TABLE_REPORT_LABELS = {
    "personnel_summary": "主要人员汇总表",
    "person_resume": "主要人员简历表",
    "person_resume_project": "主要人员简历表",
}

CERTIFICATE_PRIORITY = [
    "信息系统",
    "系统集成",
    "软件",
    "数据",
    "项目管理",
    "PMP",
    "CS",
    "ITSS",
    "DCMM",
    "云",
    "VMware",
    "VCP",
    "ACP",
    "NCP",
    "Microsoft",
    "Azure",
]


def parse_company_performance(parsed_text: str) -> dict[str, Any]:
    """Extract the tiny project-performance fact set required by the table MVP."""
    lines = [line.strip() for line in str(parsed_text or "").splitlines() if line.strip()]
    performance: dict[str, Any] = {"fact_type": "company_performance"}
    for field, labels in PERFORMANCE_FIELDS.items():
        value = first_labeled_value(lines, labels)
        if value:
            performance[field] = value
    performance["parsed_at"] = datetime.now().isoformat(timespec="seconds")
    performance["parser_version"] = "company_performance.simple.v1"
    performance["missing_fields"] = [field for field in PERFORMANCE_TABLE_FIELDS if not performance.get(field)]
    return performance


def fill_special_tables(
    template_path: str | Path,
    output_path: str | Path,
    metadata_path: str | Path = METADATA_PATH,
    debug_path: str | Path | None = None,
    extract_only: bool = False,
    missing_report: bool = True,
) -> dict[str, Any]:
    template_path = Path(template_path)
    output_path = Path(output_path)
    metadata_path = Path(metadata_path)
    debug_path = Path(debug_path) if debug_path else output_path.parent / "special_table_fill_debug.json"

    metadata = MetadataStore(metadata_path).load()
    performances = select_company_performances(metadata)
    personnel = select_personnel(metadata, limit=2)
    document = Document(str(template_path))

    debug: dict[str, Any] = {
        "template": str(template_path),
        "output": str(output_path),
        "metadata_path": str(metadata_path),
        "extract_only": extract_only,
        "missing_report": missing_report,
        "found_tables": {"performance": 0, "personnel_summary": 0, "person_resume": 0},
        "used_performances": [],
        "used_personnel": [],
        "filled_fields": [],
        "missing_fields": [],
        "default_fields": [],
        "default_used_fields": [],
        "certificate_sources": [],
        "period_normalizations": [],
        "skipped_long_role_text": [],
        "warnings": [],
    }

    contexts = table_contexts(document)
    performance_tables = [item["table"] for item in contexts if is_performance_table(item["text"])]
    summary_tables = [item["table"] for item in contexts if is_summary_table(item["text"])]
    resume_tables = [item["table"] for item in contexts if is_resume_table(item["text"])]

    debug["found_tables"] = {
        "performance": len(performance_tables),
        "personnel_summary": len(summary_tables),
        "person_resume": len(resume_tables),
    }

    if not performances:
        debug["warnings"].append("no company_performance facts found in metadata")
    if len(personnel) < 2:
        debug["warnings"].append(f"only {len(personnel)} personnel resume facts found in metadata")

    if performance_tables and performances:
        fill_performance_table(performance_tables[0], performances[0], debug)
    if summary_tables and personnel:
        fill_summary_table(summary_tables[0], personnel[:2], debug)
    if resume_tables and personnel:
        fill_resume_tables(resume_tables[0], personnel[:2], debug)

    normalize_document_tables(document)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if extract_only:
        save_special_table_fragment(document, output_path)
    else:
        document.save(str(output_path))
    debug["debug_path"] = str(debug_path)
    if missing_report:
        debug["missing_report_paths"] = write_missing_reports(output_path, debug, personnel[:2])
    debug_path.parent.mkdir(parents=True, exist_ok=True)
    debug_path.write_text(json.dumps(debug, ensure_ascii=False, indent=2), encoding="utf-8")
    return debug


def save_special_table_fragment(document: Document, output_path: Path) -> None:
    keep_only_special_table_blocks(document)
    document.save(str(output_path))


def keep_only_special_table_blocks(document: Document) -> None:
    body = document.element.body
    blocks = [child for child in list(body) if not child.tag.endswith("sectPr")]
    target_tables = target_table_elements(document)
    if not target_tables:
        return

    target_indexes = [index for index, block in enumerate(blocks) if isinstance(block, CT_Tbl) and id(block) in target_tables]
    if not target_indexes:
        return

    start = backtrack_special_region_start(blocks, min(target_indexes))
    end = extend_special_region_end(blocks, max(target_indexes))
    kept_blank = False
    for index, block in enumerate(blocks):
        keep = False
        if start <= index <= end:
            if isinstance(block, CT_Tbl):
                keep = id(block) in target_tables
                kept_blank = False
            elif isinstance(block, CT_P):
                text = paragraph_element_text(block)
                if special_related_paragraph(text):
                    keep = True
                    kept_blank = False
                elif not text.strip() and not kept_blank and nearby_target_block(blocks, index, target_tables, start, end):
                    keep = True
                    kept_blank = True
        if not keep:
            body.remove(block)


def target_table_elements(document: Document) -> set[int]:
    targets: set[int] = set()
    for item in table_contexts(document):
        text = item["text"]
        table = item["table"]
        if is_performance_table(text) or is_summary_table(text) or is_resume_table(text):
            targets.add(id(table._tbl))
    return targets


def backtrack_special_region_start(blocks: list[Any], first_table_index: int) -> int:
    start = first_table_index
    for index in range(first_table_index - 1, -1, -1):
        block = blocks[index]
        if not isinstance(block, CT_P):
            break
        text = paragraph_element_text(block)
        if special_related_paragraph(text) or not text.strip():
            start = index
            continue
        break
    return start


def extend_special_region_end(blocks: list[Any], last_table_index: int) -> int:
    end = last_table_index
    for index in range(last_table_index + 1, len(blocks)):
        block = blocks[index]
        if not isinstance(block, CT_P):
            break
        text = paragraph_element_text(block)
        if special_related_paragraph(text) or not text.strip():
            end = index
            continue
        break
    return end


def nearby_target_block(blocks: list[Any], index: int, target_tables: set[int], start: int, end: int) -> bool:
    previous_is_target = index > start and isinstance(blocks[index - 1], CT_Tbl) and id(blocks[index - 1]) in target_tables
    next_is_target = index < end and isinstance(blocks[index + 1], CT_Tbl) and id(blocks[index + 1]) in target_tables
    previous_is_related = index > start and isinstance(blocks[index - 1], CT_P) and special_related_paragraph(paragraph_element_text(blocks[index - 1]))
    next_is_related = index < end and isinstance(blocks[index + 1], CT_P) and special_related_paragraph(paragraph_element_text(blocks[index + 1]))
    return previous_is_target or next_is_target or (previous_is_related and next_is_related)


def special_related_paragraph(text: str) -> bool:
    stripped = text.strip()
    if not stripped:
        return False
    keywords = [
        "业绩文件",
        "近年完成的类似项目情况表",
        "项目团队情况",
        "拟委任的主要人员汇总表",
        "主要人员简历表",
        "表1",
        "表2",
        "注：",
        "注:",
        "层级",
        "已完项目情况证明文件",
        "业绩证明文件",
    ]
    return any(keyword in stripped for keyword in keywords)


def paragraph_element_text(paragraph: CT_P) -> str:
    return "".join(node.text or "" for node in paragraph.iter() if node.tag.endswith("}t"))


def write_missing_reports(output_path: Path, debug: dict[str, Any], personnel: list[dict[str, Any]]) -> dict[str, str]:
    report = build_missing_report(debug, personnel)
    md_path = output_path.with_name(f"{output_path.stem}.missing_report.md")
    json_path = output_path.with_name(f"{output_path.stem}.missing_report.json")
    md_path.write_text(render_missing_report_markdown(report), encoding="utf-8")
    json_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return {"markdown": str(md_path), "json": str(json_path)}


def build_missing_report(debug: dict[str, Any], personnel: list[dict[str, Any]]) -> dict[str, Any]:
    missing_by_person: dict[str, list[dict[str, Any]]] = {}
    for item in list(debug.get("missing_fields") or []):
        person_name = text_value(item.get("person_name"))
        if not person_name:
            continue
        missing_by_person.setdefault(person_name, []).append(normalize_missing_item(item))

    people = []
    for person in personnel:
        person_name = text_value(person.get("person_name")) or "待确认"
        summary = person.get("_special_summary_values") if isinstance(person.get("_special_summary_values"), dict) else {}
        missing_items = dedupe_missing_items(missing_by_person.get(person_name, []))
        questions = sorted_unique([item["question"] for item in missing_items if item.get("question")])
        people.append(
            {
                "person_name": person_name,
                "department": person_department(person),
                "bid_level": text_value(summary.get("level")) or "待确认",
                "bid_job": text_value(summary.get("job")) or "待确认",
                "missing_fields": missing_items,
                "questions": questions,
            }
        )

    departments: dict[str, list[dict[str, Any]]] = {}
    for person in people:
        departments.setdefault(person["department"], []).append(person)
    return {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "source_output": debug.get("output"),
        "departments": [{"department": department, "people": people} for department, people in departments.items()],
    }


def normalize_missing_item(item: dict[str, Any]) -> dict[str, Any]:
    field = text_value(item.get("field"))
    table = text_value(item.get("table"))
    return {
        "field": field,
        "field_label": FIELD_REPORT_LABELS.get(field, field or "待确认"),
        "table": table,
        "table_label": TABLE_REPORT_LABELS.get(table, table or "待确认"),
        "reason": text_value(item.get("reason")) or "source_missing",
        "usage": FIELD_USAGE.get(field, "用于专项表格填报"),
        "question": FIELD_QUESTIONS.get(field, f"请补充{FIELD_REPORT_LABELS.get(field, field)}。"),
    }


def dedupe_missing_items(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    result = []
    seen: set[tuple[str, str]] = set()
    for item in items:
        key = (text_value(item.get("table_label")), text_value(item.get("field_label")))
        if key in seen:
            continue
        seen.add(key)
        result.append(item)
    return result


def render_missing_report_markdown(report: dict[str, Any]) -> str:
    lines = ["# 专项表格缺失信息清单", ""]
    lines.append(f"生成时间：{report.get('generated_at', '')}")
    lines.append("")
    for group in list(report.get("departments") or []):
        lines.append(f"## 部门：{group.get('department') or '待确认'}")
        lines.append("")
        for person in list(group.get("people") or []):
            lines.append(f"### 人员：{person.get('person_name') or '待确认'}")
            lines.append(f"- 部门：{person.get('department') or '待确认'}")
            lines.append(f"- 本次投标层级：{person.get('bid_level') or '待确认'}")
            lines.append(f"- 本次投标职务：{person.get('bid_job') or '待确认'}")
            lines.append("")
            lines.append("缺失字段：")
            if person.get("missing_fields"):
                for item in person["missing_fields"]:
                    lines.append(
                        f"- {item['field_label']}：{item['usage']}；所在表格：{item['table_label']}；原因：{item['reason']}"
                    )
            else:
                lines.append("- 暂无")
            lines.append("")
            lines.append("建议向部门确认：")
            if person.get("questions"):
                for question in person["questions"]:
                    lines.append(f"- {question}")
            else:
                lines.append("- 暂无")
            lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def person_department(person: dict[str, Any]) -> str:
    resume = person.get("resume") if isinstance(person.get("resume"), dict) else {}
    profile = person.get("personnel_profile") if isinstance(person.get("personnel_profile"), dict) else {}
    for value in [resume.get("department"), profile.get("department")]:
        text = text_value(value)
        if text:
            return text
    source_text = " ".join(
        text_value(value)
        for value in [
            resume.get("source_folder"),
            person.get("source_folder"),
            " ".join(str(tag) for tag in list(person.get("tags") or [])),
        ]
        if value
    )
    match = re.search(r"[\u4e00-\u9fffA-Za-z0-9]+(?:事业部|部门|中心|部)", source_text)
    return match.group(0) if match else "待确认"


def sorted_unique(values: list[str]) -> list[str]:
    result = []
    seen: set[str] = set()
    for value in values:
        text = text_value(value)
        if text and text not in seen:
            seen.add(text)
            result.append(text)
    return result


def fill_performance_table(table: Table, performance: dict[str, Any], debug: dict[str, Any]) -> None:
    debug["used_performances"].append(
        {
            "source_file_id": performance.get("source_file_id"),
            "original_name": performance.get("original_name"),
            "project_name": performance.get("project_name"),
        }
    )
    header_indexes = header_indexes_for(table, {field: PERFORMANCE_FIELDS[field] for field in PERFORMANCE_TABLE_FIELDS})
    if header_indexes:
        header_row_index = int(header_indexes.pop("__row_index", 0))
        row = first_empty_or_new_row(table, start_index=header_row_index + 1)
        for field in PERFORMANCE_TABLE_FIELDS:
            value = text_value(performance.get(field))
            col = header_indexes.get(field)
            if col is None:
                debug["missing_fields"].append({"table": "performance", "field": field, "reason": "column_not_found"})
                continue
            write_cell(row.cells[col], value)
            record_fill(debug, "performance", field, value)
        return

    for field in PERFORMANCE_TABLE_FIELDS:
        value = text_value(performance.get(field))
        if fill_label_neighbor(table, PERFORMANCE_FIELDS[field], value):
            record_fill(debug, "performance", field, value)
        else:
            debug["missing_fields"].append({"table": "performance", "field": field, "reason": "label_not_found"})


def fill_summary_table(table: Table, personnel: list[dict[str, Any]], debug: dict[str, Any]) -> None:
    header_indexes = header_indexes_for(table, SUMMARY_LABELS)
    if not header_indexes:
        debug["warnings"].append("personnel summary table header not found")
        return
    data_start = int(header_indexes.pop("__row_index", 0)) + 1
    for index, person in enumerate(personnel):
        row = row_at_or_new(table, data_start + index)
        values = build_person_summary(person, index, debug)
        missing_reasons = values.pop("__missing_reasons", {})
        values.pop("__certificate_source", None)
        for field in SUMMARY_FIELDS:
            col = header_indexes.get(field)
            if col is None:
                debug["missing_fields"].append({"table": "personnel_summary", "field": field, "reason": "column_not_found"})
                continue
            value = values.get(field, "")
            write_cell(row.cells[col], value, missing=is_missing_value(value))
            record_fill(debug, "personnel_summary", field, value, person, missing_reasons.get(field))
    debug["used_personnel"] = [
        {"file_id": person.get("file_id"), "original_name": person.get("original_name"), "person_name": person.get("person_name")}
        for person in personnel
    ]


def fill_resume_tables(template_table: Table, personnel: list[dict[str, Any]], debug: dict[str, Any]) -> None:
    tables = [template_table]
    while len(tables) < len(personnel):
        tables.append(clone_resume_section_after(tables[-1], template_table))

    for index, person in enumerate(personnel):
        values = build_person_resume(person, index, debug)
        missing_reasons = values.pop("__missing_reasons", {})
        table = tables[index]
        for field, labels in RESUME_LABELS.items():
            value = values.get(field, "")
            if fill_label_neighbor(table, labels, value, missing=is_missing_value(value)):
                record_fill(debug, "person_resume", field, value, person, missing_reasons.get(field))
            else:
                debug["missing_fields"].append({"table": "person_resume", "field": field, "reason": "label_not_found"})
        fill_resume_project_rows(table, values.get("projects", []), person, debug)


def fill_resume_project_rows(table: Table, projects: list[dict[str, str]], person: dict[str, Any], debug: dict[str, Any]) -> None:
    header_row_index, indexes = find_project_header(table)
    if header_row_index is None or not indexes:
        if not projects:
            debug["missing_fields"].append({"table": "person_resume", "field": "projects", "person_name": person.get("person_name")})
        return
    if not projects:
        debug["missing_fields"].append({"table": "person_resume", "field": "projects", "person_name": person.get("person_name")})
        return
    data_start = header_row_index + 1
    for offset, project in enumerate(projects[:3]):
        row = row_at_or_new(table, data_start + offset)
        missing_reasons = project.get("__missing_reasons", {})
        default_used = set(project.get("__default_used", []))
        for field, labels in PROJECT_LABELS.items():
            col = indexes.get(field)
            if col is None:
                continue
            value = text_value(project.get(field))
            write_cell(row.cells[col], value, missing=is_missing_value(value))
            if field in default_used:
                debug["default_used_fields"].append(
                    {
                        "table": "person_resume_project",
                        "field": field,
                        "person_name": person.get("person_name"),
                        "value": value,
                        "reason": "project_role_missing_used_bid_job",
                    }
                )
            record_fill(debug, "person_resume_project", field, value, person, missing_reasons.get(field))


def build_person_summary(person: dict[str, Any], index: int, debug: dict[str, Any]) -> dict[str, str]:
    if isinstance(person.get("_special_summary_values"), dict):
        return copy.deepcopy(person["_special_summary_values"])
    resume = person.get("resume") or {}
    profile = person.get("personnel_profile") or {}
    certificate, certificate_source = choose_certificate(resume, person)
    debug["certificate_sources"].append(certificate_source)
    level = first_text(resume, ["level", "project_role", "role"]) or first_role_from_profile(profile)
    if not level:
        level = "项目负责人" if index == 0 else "其他主要人员"
        default_item = {"table": "personnel_summary", "field": "level", "person_name": person.get("person_name"), "value": level, "reason": "source_missing"}
        debug["default_fields"].append(default_item)
        debug["default_used_fields"].append(default_item)
    job = first_text(resume, ["job", "job_title", "position", "post"]) or first_work_position(resume)
    certificate_name = certificate_display_name(certificate)
    certificate_major = certificate_major_value(certificate, resume)
    values = {
        "level": level,
        "job": value_or_missing(job),
        "name": value_or_missing(person.get("person_name")),
        "title": value_or_missing(first_text(resume, ["title", "professional_title", "technical_title", "职称"])),
        "certificate_name": value_or_missing(certificate_name),
        "certificate_level": value_or_missing(certificate.get("level") or certificate.get("certificate_level")),
        "certificate_no": value_or_missing(certificate.get("certificate_no") or certificate.get("number") or certificate.get("certificate_number")),
        "certificate_major": value_or_missing(certificate_major),
        "remark": "",
        "__certificate_source": certificate_source,
        "__missing_reasons": {},
    }
    for field in ["job", "title", "certificate_name", "certificate_level", "certificate_no", "certificate_major"]:
        if is_missing_value(values[field]):
            values["__missing_reasons"][field] = certificate_source.get("missing_reason") if field.startswith("certificate") and not certificate else "source_missing"
    if is_missing_value(values["name"]):
        values["__missing_reasons"]["name"] = "not_in_resume"
    person["_special_summary_values"] = copy.deepcopy(values)
    return values


def build_person_resume(person: dict[str, Any], index: int, debug: dict[str, Any]) -> dict[str, Any]:
    summary = build_person_summary(person, index, debug)
    resume = person.get("resume") or {}
    start_work_date = first_text(resume, ["start_work_date", "work_start_date", "first_work_date", "参加工作时间"])
    work_years = resolve_work_years(resume, start_work_date)
    values = {
        "level": summary["level"],
        "name": summary["name"],
        "gender": value_or_missing(resume.get("gender")),
        "age": value_or_missing(resolve_age(resume)),
        "education": value_or_missing(first_text(resume, ["education", "degree", "education_background", "学历"])),
        "start_work_date": value_or_missing(start_work_date),
        "work_years": value_or_missing(work_years),
        "projects": normalize_resume_projects(resume, summary["job"], person, debug),
        "__missing_reasons": {},
    }
    for field in ["gender", "age", "education", "start_work_date", "work_years"]:
        if is_missing_value(values[field]):
            values["__missing_reasons"][field] = "not_in_resume"
    values["__missing_reasons"].update(summary.get("__missing_reasons", {}))
    return values


def select_company_performances(metadata: dict[str, Any]) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    seen: set[str] = set()
    top = metadata.get("company_performances")
    if isinstance(top, dict):
        for file_id, value in top.items():
            if isinstance(value, dict):
                item = dict(value)
                item.setdefault("source_file_id", file_id)
                add_performance(items, seen, item)
    elif isinstance(top, list):
        for value in top:
            if isinstance(value, dict):
                add_performance(items, seen, value)

    for record in list(metadata.get("files") or []):
        value = record.get("company_performance")
        if isinstance(value, dict):
            item = dict(value)
            item.setdefault("source_file_id", record.get("id"))
            item.setdefault("original_name", record.get("original_name"))
            add_performance(items, seen, item)
    return [item for item in items if item.get("project_name")]


def add_performance(items: list[dict[str, Any]], seen: set[str], item: dict[str, Any]) -> None:
    key = str(item.get("source_file_id") or item.get("project_name") or len(items))
    if key in seen:
        return
    seen.add(key)
    items.append(item)


def select_personnel(metadata: dict[str, Any], limit: int = 2) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    seen: set[str] = set()
    for record in list(metadata.get("files") or []):
        resume = resume_from_record(record)
        if not resume:
            continue
        person_name = text_value(resume.get("person_name") or resume.get("name"))
        if not person_name:
            continue
        file_id = str(record.get("id") or "")
        if file_id in seen:
            continue
        seen.add(file_id)
        result.append(
            {
                "file_id": record.get("id"),
                "original_name": record.get("original_name"),
                "person_name": person_name,
                "resume": resume,
                "personnel_profile": record.get("personnel_profile") if isinstance(record.get("personnel_profile"), dict) else {},
            }
        )
        if len(result) >= limit:
            return result

    top = metadata.get("resumes")
    if isinstance(top, dict):
        for file_id, resume in top.items():
            if file_id in seen or not isinstance(resume, dict):
                continue
            person_name = text_value(resume.get("person_name") or resume.get("name"))
            if not person_name:
                continue
            result.append({"file_id": file_id, "original_name": "", "person_name": person_name, "resume": resume, "personnel_profile": {}})
            if len(result) >= limit:
                break
    return result


def resume_from_record(record: dict[str, Any]) -> dict[str, Any]:
    if isinstance(record.get("resume_parse"), dict):
        return record["resume_parse"]
    structured = record.get("structured_facts")
    if isinstance(structured, dict) and isinstance(structured.get("resume_parse"), dict):
        return structured["resume_parse"]
    return {}


def table_contexts(document: Document) -> list[dict[str, Any]]:
    contexts = []
    previous_paragraphs: list[str] = []
    for block in iter_blocks(document):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if text:
                previous_paragraphs.append(text)
                previous_paragraphs = previous_paragraphs[-6:]
            continue
        table_text = table_plain_text(block)
        context_text = "\n".join([*previous_paragraphs, table_text])
        contexts.append({"table": block, "text": context_text})
    return contexts


def iter_blocks(document: Document) -> Iterable[Paragraph | Table]:
    parent = document.element.body
    for child in parent.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def is_performance_table(text: str) -> bool:
    labels = ["项目法人", "项目名称", "项目规模", "服务范围", "服务期间", "项目目前进度", "服务人员人数", "建设单位联系人及电话"]
    hits = count_contains(text, labels)
    return hits >= 5 or ("近年完成的类似项目情况表" in text and hits >= 3)


def is_summary_table(text: str) -> bool:
    labels = ["层级", "职务", "姓名", "职称", "执业或职业资格证明", "证书名称", "级别", "证号", "专业", "备注"]
    hits = count_contains(text, labels)
    return hits >= 8 or ("拟委任的主要人员汇总表" in text and hits >= 5)


def is_resume_table(text: str) -> bool:
    labels = ["主要人员简历表", "层级", "姓名", "性别", "年龄", "学历", "参加工作时间", "工作年限", "主要资历、经验及承担过的项目", "项目名称", "建设规模", "项目起止时间", "担负的技术职务", "获奖情况"]
    hits = count_contains(text, labels)
    return hits >= 9 or ("主要人员简历表" in text and hits >= 6)


def count_contains(text: str, labels: list[str]) -> int:
    normalized = normalize_label(text)
    return sum(1 for label in labels if normalize_label(label) in normalized)


def table_plain_text(table: Table) -> str:
    return "\n".join("\t".join(cell.text.strip() for cell in row.cells) for row in table.rows)


def header_indexes_for(table: Table, labels: dict[str, list[str]]) -> dict[str, int]:
    best: dict[str, int] = {}
    best_count = 0
    for row_index, row in enumerate(table.rows[:4]):
        found: dict[str, int] = {"row_index": row_index}
        for field, candidates in labels.items():
            for col_index, cell in enumerate(row.cells):
                if any(label_matches(cell.text, label) for label in candidates):
                    found[field] = col_index
                    break
        count = len([key for key in found if key != "row_index"])
        if count > best_count:
            best = found
            best_count = count
    if "row_index" in best:
        best["__row_index"] = best.pop("row_index")
    return best if best_count >= 2 else {}


def find_project_header(table: Table) -> tuple[int | None, dict[str, int]]:
    best_row = None
    best: dict[str, int] = {}
    for row_index, row in enumerate(table.rows):
        found: dict[str, int] = {}
        for field, labels in PROJECT_LABELS.items():
            for col_index, cell in enumerate(row.cells):
                if any(label_matches(cell.text, label) for label in labels):
                    found[field] = col_index
                    break
        if len(found) > len(best):
            best_row = row_index
            best = found
    return (best_row, best) if len(best) >= 3 else (None, {})


def fill_label_neighbor(table: Table, labels: list[str], value: str, missing: bool = False) -> bool:
    for row in table.rows:
        cells = list(row.cells)
        for index, cell in enumerate(cells):
            if not any(label_matches(cell.text, label) for label in labels):
                continue
            for target_index in range(index + 1, len(cells)):
                target = cells[target_index]
                if target._tc is cell._tc:
                    continue
                if any(label_matches(target.text, label) for label in labels):
                    continue
                write_cell(target, value, missing=missing)
                return True
    return False


def first_empty_or_new_row(table: Table, start_index: int = 1):
    for row in table.rows[start_index:]:
        if not any(cell.text.strip() for cell in row.cells):
            return row
    return clone_row_after(table, len(table.rows) - 1)


def row_at_or_new(table: Table, index: int):
    while len(table.rows) <= index:
        clone_row_after(table, len(table.rows) - 1)
    return table.rows[index]


def clone_row_after(table: Table, row_index: int):
    source = table.rows[row_index]
    new_tr = copy.deepcopy(source._tr)
    source._tr.addnext(new_tr)
    new_row = table.rows[row_index + 1]
    for cell in unique_cells(new_row.cells):
        write_cell(cell, "")
    return new_row


def clone_resume_section_after(anchor_table: Table, template_table: Table) -> Table:
    anchor = anchor_table._tbl
    blank = OxmlElement("w:p")
    anchor.addnext(blank)
    anchor = blank
    new_tbl = copy.deepcopy(template_table._tbl)
    anchor.addnext(new_tbl)
    cloned = Table(new_tbl, template_table._parent)
    clear_resume_table_values(cloned)
    return cloned


def clear_resume_table_values(table: Table) -> None:
    for row in table.rows:
        for cell in unique_cells(row.cells):
            if not is_headerish_cell(cell):
                write_cell(cell, "")


def clone_table_after(table: Table) -> Table:
    new_tbl = copy.deepcopy(table._tbl)
    table._tbl.addnext(new_tbl)
    cloned = Table(new_tbl, table._parent)
    for row in cloned.rows:
        for cell in unique_cells(row.cells):
            if not is_headerish_cell(cell):
                write_cell(cell, "")
    return cloned


def unique_cells(cells: Iterable[_Cell]) -> list[_Cell]:
    result = []
    seen: set[int] = set()
    for cell in cells:
        key = id(cell._tc)
        if key in seen:
            continue
        seen.add(key)
        result.append(cell)
    return result


def is_headerish_cell(cell: _Cell) -> bool:
    text = cell.text.strip()
    labels = [
        *sum(RESUME_LABELS.values(), []),
        *sum(PROJECT_LABELS.values(), []),
        "主要资历、经验及承担过的项目",
    ]
    return any(label_matches(text, label) for label in labels)


def write_cell(cell: _Cell, value: Any, missing: bool = False) -> None:
    text = text_value(value)
    if not cell.paragraphs:
        paragraph = cell.add_paragraph()
        write_runs(paragraph, text, missing=missing)
        return
    first = cell.paragraphs[0]
    for run in first.runs:
        run.text = ""
    write_runs(first, text, missing=missing)
    for paragraph in cell.paragraphs[1:]:
        for run in paragraph.runs:
            run.text = ""


def write_runs(paragraph, text: str, missing: bool = False) -> None:
    if not text:
        run = paragraph.add_run("")
        apply_run_style(run, missing=False)
        return
    if missing or MISSING_VALUE not in text:
        run = paragraph.add_run(text)
        apply_run_style(run, missing=missing)
        return
    parts = text.split(MISSING_VALUE)
    for index, part in enumerate(parts):
        if part:
            run = paragraph.add_run(part)
            apply_run_style(run, missing=False)
        if index < len(parts) - 1:
            run = paragraph.add_run(MISSING_VALUE)
            apply_run_style(run, missing=True)


def apply_run_style(run, missing: bool = False) -> None:
    run.font.size = FILL_FONT_SIZE
    if missing:
        run.font.color.rgb = MISSING_COLOR


def normalize_document_tables(document: Document) -> None:
    for table in document.tables:
        for row in table.rows:
            try:
                if row.height_rule == WD_ROW_HEIGHT_RULE.EXACTLY:
                    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            except Exception:
                continue


def first_labeled_value(lines: list[str], labels: list[str]) -> str:
    for line in lines:
        for label in labels:
            pattern = rf"^\s*{re.escape(label)}\s*[：:\t ]+\s*(?P<value>.+?)\s*$"
            match = re.match(pattern, line)
            if match:
                return match.group("value").strip()
            compact_line = normalize_label(line)
            compact_label = normalize_label(label)
            if compact_line.startswith(compact_label):
                value = line[len(label) :].strip(" ：:\t")
                if value:
                    return value
    return ""


def first_text(data: dict[str, Any], keys: list[str]) -> str:
    for key in keys:
        value = data.get(key)
        if isinstance(value, str) and value.strip():
            return value.strip()
    return ""


def first_work_position(resume: dict[str, Any]) -> str:
    for item in list(resume.get("work_history") or []):
        if isinstance(item, dict) and text_value(item.get("position")):
            return text_value(item.get("position"))
    return ""


def first_role_from_profile(profile: dict[str, Any]) -> str:
    role_tags = list(profile.get("role_tags") or [])
    if "project_manager" in role_tags:
        return "项目负责人"
    return ""


def choose_certificate(resume: dict[str, Any], person: dict[str, Any]) -> tuple[dict[str, Any], dict[str, Any]]:
    raw_certificates = [item for item in list(resume.get("certificates") or []) if isinstance(item, dict)]
    certificates = [item for item in raw_certificates if certificate_usable_for_person(item, person)]
    base_source = {
        "person_name": person.get("person_name"),
        "file_id": person.get("file_id"),
        "source": "resume.certificates",
        "raw_certificate_count": len(raw_certificates),
        "usable_certificate_count": len(certificates),
    }
    if not certificates:
        reason = "ocr_missing" if any(item.get("needs_ocr") for item in raw_certificates) else "not_in_resume"
        return {}, {**base_source, "certificate_name": "", "missing_reason": reason}

    def score(certificate: dict[str, Any]) -> int:
        text = json.dumps(certificate, ensure_ascii=False)
        return max([len(CERTIFICATE_PRIORITY) - index for index, keyword in enumerate(CERTIFICATE_PRIORITY) if keyword.lower() in text.lower()] or [0])

    certificate = sorted(certificates, key=score, reverse=True)[0]
    return certificate, {
        **base_source,
        "certificate_name": certificate_display_name(certificate),
        "certificate_source_type": certificate.get("source_type"),
        "ownership_source": certificate.get("ownership_source") or "resume_scope",
        "ownership_file_id": certificate.get("ownership_file_id"),
        "certificate_person_name": certificate.get("person_name") or certificate.get("holder_name"),
    }


def certificate_usable_for_person(certificate: dict[str, Any], person: dict[str, Any]) -> bool:
    if certificate.get("needs_ocr"):
        return False
    if certificate.get("usable") is False and not certificate_display_name(certificate):
        return False
    if not certificate_display_name(certificate):
        return False
    person_name = text_value(person.get("person_name"))
    file_id = text_value(person.get("file_id"))
    cert_person = text_value(certificate.get("person_name") or certificate.get("holder_name"))
    cert_file_id = text_value(certificate.get("ownership_file_id") or certificate.get("file_id"))
    if cert_person and person_name and cert_person != person_name:
        return False
    if cert_file_id and file_id and cert_file_id != file_id:
        return False
    return True


def certificate_display_name(certificate: dict[str, Any]) -> str:
    for key in ["display_name", "standard_name", "normalized_name", "certificate_full_name", "name"]:
        value = text_value(certificate.get(key))
        if value:
            return value
    return ""


def certificate_major_value(certificate: dict[str, Any], resume: dict[str, Any]) -> str:
    for key in ["major", "professional", "profession", "专业"]:
        value = text_value(certificate.get(key))
        if value:
            return value
    return first_text(resume, ["education_major", "major", "student_major", "所学专业"])


def resolve_age(resume: dict[str, Any]) -> str:
    age = text_value(resume.get("age"))
    if age:
        return age
    birthday = first_text(resume, ["birthday", "birth_date", "date_of_birth", "出生日期"])
    parsed = parse_date_value(birthday)
    if not parsed:
        return ""
    today = date.today()
    years = today.year - parsed.year - ((today.month, today.day) < (parsed.month, parsed.day))
    return str(years) if years >= 0 else ""


def resolve_work_years(resume: dict[str, Any], start_work_date: str) -> str:
    parsed = parse_date_value(start_work_date)
    if parsed:
        today = date.today()
        years = today.year - parsed.year - ((today.month, today.day) < (parsed.month, parsed.day))
        return str(max(years, 0))
    return text_value(resume.get("work_years"))


def parse_date_value(value: Any) -> date | None:
    text = text_value(value)
    if not text:
        return None
    match = re.search(r"(?P<year>19\d{2}|20\d{2})(?:[年./-]\s*(?P<month>\d{1,2}))?(?:[月./-]\s*(?P<day>\d{1,2}))?", text)
    if not match:
        return None
    year = int(match.group("year"))
    month = int(match.group("month") or 1)
    day = int(match.group("day") or 1)
    try:
        return date(year, month, day)
    except ValueError:
        return None


def normalize_resume_projects(resume: dict[str, Any], default_job: str, person: dict[str, Any], debug: dict[str, Any]) -> list[dict[str, Any]]:
    projects = []
    project_items = resume.get("project_experiences") or resume.get("projects") or resume.get("project_experience") or []
    for item in list(project_items):
        if not isinstance(item, dict):
            continue
        role, role_meta = project_role_value(item, default_job, person, debug)
        default_used: list[str] = []
        if role_meta.get("default_used"):
            default_used.append("role")
        award = award_value(item)
        raw_period = project_period(item)
        period_result = normalize_project_period(raw_period)
        period_debug = {
            "table": "person_resume_project",
            "field": "date",
            "person_name": person.get("person_name"),
            "project_name": text_value(item.get("project_name") or item.get("name")),
            **period_result["debug"],
        }
        debug["period_normalizations"].append(period_debug)
        project = {
            "name": value_or_missing(item.get("project_name") or item.get("name")),
            "scale": value_or_missing(item.get("project_scale") or item.get("scale") or item.get("construction_scale")),
            "date": value_or_missing(period_result["value"]),
            "role": value_or_missing(role),
            "award": value_or_missing(award),
            "__missing_reasons": {},
            "__default_used": default_used,
        }
        for field in ["name", "scale", "date", "role", "award"]:
            if is_missing_value(project[field]):
                project["__missing_reasons"][field] = "not_in_resume"
        if period_result["debug"].get("date_precision") == "partial":
            project["__missing_reasons"]["date"] = "source_missing"
        projects.append(
            project
        )
    return [project for project in projects if any(project.values())][:3]


def project_period(item: dict[str, Any]) -> str:
    start = text_value(item.get("start_date"))
    end = text_value(item.get("end_date"))
    if start or end:
        return f"{start} 至 {end}".strip()
    return text_value(item.get("period") or item.get("date") or item.get("time") or item.get("project_period"))


def normalize_project_period(value: Any) -> dict[str, Any]:
    original = text_value(value)
    if not original:
        return {
            "value": "",
            "debug": {"original_period": "", "normalized_period": MISSING_VALUE, "date_precision": "missing"},
        }
    values = extract_year_month_values(original)
    if len(values) >= 2:
        start = format_year_month_value(values[0])
        end = format_year_month_value(values[1])
        if values[0]["month"] is None and values[1]["month"] is None:
            normalized = f"{values[0]['year']}年-{values[1]['year']}年"
            precision = "year_only"
        else:
            normalized = f"{start}-{end}"
            precision = "year_month"
        return {
            "value": normalized,
            "debug": {"original_period": original, "normalized_period": normalized, "date_precision": precision},
        }
    if len(values) == 1:
        only = format_year_month_value(values[0])
        if values[0]["month"] is None:
            normalized = only
            precision = "year_only"
        elif re.search(r"至|到|-|—|~", original):
            normalized = f"{only}-{MISSING_VALUE}"
            precision = "partial"
        else:
            normalized = only
            precision = "year_month"
        return {
            "value": normalized,
            "debug": {"original_period": original, "normalized_period": normalized, "date_precision": precision},
        }
    return {
        "value": original,
        "debug": {"original_period": original, "normalized_period": original, "date_precision": "raw"},
    }


def extract_year_month_values(value: str) -> list[dict[str, int | None]]:
    results: list[dict[str, int | None]] = []
    pattern = re.compile(r"(?P<year>19\d{2}|20\d{2})(?:\s*[年./-]\s*(?P<month>\d{1,2})\s*月?)?")
    for match in pattern.finditer(value):
        year = int(match.group("year"))
        month_text = match.group("month")
        month = int(month_text) if month_text else None
        if month is not None and not 1 <= month <= 12:
            month = None
        results.append({"year": year, "month": month})
    return results


def format_year_month_value(value: dict[str, int | None]) -> str:
    year = int(value["year"])
    month = value.get("month")
    if month is None:
        return f"{year}年"
    return f"{year}年{int(month):02d}月"


def project_role_value(
    item: dict[str, Any],
    default_job: str,
    person: dict[str, Any],
    debug: dict[str, Any],
) -> tuple[str, dict[str, Any]]:
    for key in ["technical_role", "role", "position", "duty_role"]:
        candidate = text_value(item.get(key))
        if not candidate:
            continue
        if looks_like_long_role_description(candidate):
            debug["skipped_long_role_text"].append(
                {
                    "person_name": person.get("person_name"),
                    "project_name": text_value(item.get("project_name") or item.get("name")),
                    "field": key,
                    "value": candidate,
                    "reason": "looks_like_responsibility_description",
                }
            )
            continue
        return candidate, {"source": key}
    if default_job and not is_missing_value(default_job):
        return default_job, {"source": "bid_job", "default_used": True}
    return "", {"source": "missing"}


def looks_like_long_role_description(value: str) -> bool:
    text = text_value(value)
    chinese_chars = re.findall(r"[\u4e00-\u9fff]", text)
    punctuation_count = len(re.findall(r"[，,。、；;]", text))
    return len(chinese_chars) > 20 or punctuation_count >= 2


def award_value(item: dict[str, Any]) -> str:
    for key in ["awards", "honors", "award"]:
        if key not in item:
            continue
        value = item.get(key)
        if isinstance(value, list):
            return "无" if not value else "；".join(text_value(part) for part in value if text_value(part))
        text = text_value(value)
        if text:
            return text
        return "无"
    return ""


def record_fill(
    debug: dict[str, Any],
    table: str,
    field: str,
    value: str,
    person: dict[str, Any] | None = None,
    missing_reason: str | None = None,
) -> None:
    item = {"table": table, "field": field, "filled": bool(text_value(value))}
    if person:
        item["person_name"] = person.get("person_name")
        item["file_id"] = person.get("file_id")
    if MISSING_VALUE in text_value(value):
        debug["missing_fields"].append({**item, "filled": False, "reason": missing_reason or "source_missing", "value": MISSING_VALUE})
    elif text_value(value):
        debug["filled_fields"].append(item)
    else:
        debug["missing_fields"].append({**item, "reason": missing_reason or "source_value_empty"})


def label_matches(text: str, label: str) -> bool:
    return normalize_label(label) in normalize_label(text)


def normalize_label(value: Any) -> str:
    text = str(value or "")
    return re.sub(r"[\s　（）()：:；;、/|\\\[\]【】]+", "", text)


def text_value(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return text if text.lower() != "none" else ""


def value_or_missing(value: Any) -> str:
    text = text_value(value)
    return text if text else MISSING_VALUE


def is_missing_value(value: Any) -> bool:
    return text_value(value) == MISSING_VALUE
